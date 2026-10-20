"""Generate slide decks with Claude and render them to .pptx."""

from __future__ import annotations

import io
import re
from typing import Any

from services.claude_client import generate_json
from services.presentation_images import IMAGE_ID_RE, attach_images, image_path, image_url
from services.presentation_themes import DEFAULT_THEME, THEMES, Theme, accent_split, get_theme

MIN_SLIDES = 5
MAX_SLIDES = 15
MAX_TOPIC_CHARS = 1000
MAX_DETAILS_CHARS = 8000
MAX_BULLETS = 6
MAX_BULLET_CHARS = 220
MAX_TITLE_CHARS = 120
MAX_NOTES_CHARS = 2400
MAX_PRESENTER_CHARS = 160
_SLIDE_COUNT_RE = re.compile(r"\b(\d{1,2})\s*[- ]?\s*(?:slides?|слайд\w*)\b", re.IGNORECASE)
_DURATION_RE = re.compile(
    r"\b(\d{1,2})(?:\s*[-–—]\s*(\d{1,2}))?\s*[- ]?\s*(?:minutes?|mins?|минут\w*|мин)\b", re.IGNORECASE
)
# People speak ~130 words/min, but the model overshoots requested note lengths by ~25%.
_NOTES_WORDS_PER_MINUTE = 105
MAX_ITEMS = 5
MAX_ITEM_HEADING_CHARS = 60
MAX_ITEM_TEXT_CHARS = 180
MAX_QUERY_CHARS = 80

LAYOUTS = ("bullets", "image", "cards", "timeline", "two_column", "section", "closing")
PHOTO_LAYOUTS = ("image", "section")
ITEM_LAYOUTS = {"cards": 4, "timeline": MAX_ITEMS}

TONES = {
    "academic": "an academic audience (university seminar or lecture); precise, evidence-based wording",
    "business": "a professional business audience; concise, outcome-focused wording",
    "simple": "a general audience; plain, easy-to-follow wording with no jargon",
}

LANGUAGES = {
    "auto": "the same language as the topic",
    "en": "English",
    "ru": "Russian",
    "uz": "Uzbek",
    "kk": "Kazakh",
    "es": "Spanish",
    "de": "German",
    "fr": "French",
}

_SYSTEM_PROMPT = """You are an expert presentation designer who builds clear, visual, well-structured slide decks.
Return ONLY a JSON object, no markdown fences, matching this schema:
{
  "title": "deck title",
  "subtitle": "one-line subtitle",
  "presenter": "presenter name and organization exactly as the user gave them, joined by ' · '; empty if not given",
  "image_query": "cover photo search query (see rules)",
  "slides": [
    {
      "layout": "bullets" | "image" | "cards" | "timeline" | "two_column" | "section" | "closing",
      "title": "slide title",
      "subtitle": "only for section/closing slides, optional",
      "bullets": ["bullets/image/closing layouts: 3-5 items (image: 3-4)"],
      "items": [{"heading": "2-4 words", "text": "one short sentence"}],
      "left": {"heading": "two_column only", "bullets": ["2-4 items"]},
      "right": {"heading": "two_column only", "bullets": ["2-4 items"]},
      "image_query": "image and section layouts only (see rules)",
      "notes": "speaker notes: the words the presenter says on this slide (see rules)"
    }
  ]
}
Layouts:
- "image": 3-4 short bullets next to a large photo. Use it for about a third of the content slides.
- "cards": 3-4 parallel key ideas (factors, benefits, principles) in "items".
- "timeline": 3-5 chronological stages, steps, or historical milestones in "items", in order.
- "two_column": comparisons (pros/cons, before/after, theory/practice).
- "bullets": plain list, for dense content that fits no visual layout.
- "section": divides a long deck into parts; use sparingly.
- "closing": the last slide (conclusion / key takeaways).
Rules:
- Make the deck feel designed: vary layouts and never use the same layout more than twice in a row.
- Bullets and item texts are short phrases (max ~15 words), never full paragraphs.
- Follow every explicit requirement in the user's request (title, presenter, audience, purpose, structure).
- Notes are a script spoken in the first person, never instructions to the presenter ("Provide contact details",
  "Emphasize that..."). If the user gives a talk duration, the notes across all slides must fill it at about
  130 spoken words per minute; otherwise write 2-4 sentences per slide.
- image_query is 2-4 English words naming a literal, photographable subject, e.g. "students working on laptops",
  "wind turbines", "doctor examining patient", "university library". Name people, places, or objects a camera can
  capture. No abstract words (concept, impact, future, network, digital, global, transformation), no text,
  charts, or diagrams. Each query in the deck must be different.
- Facts must be accurate; do not invent statistics, quotes, or citations. Only use specific numbers or
  percentages that are well established or come from the user's material."""


class PresentationError(Exception):
    """Raised when a deck cannot be generated or rendered."""


def _clean(value: Any, limit: int) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    return text[:limit].rstrip()


def _clean_bullets(value: Any, limit: int = MAX_BULLETS) -> list[str]:
    if not isinstance(value, list):
        return []
    bullets = [_clean(item, MAX_BULLET_CHARS) for item in value]
    return [b for b in bullets if b][:limit]


def _clean_column(value: Any) -> dict[str, Any]:
    col = value if isinstance(value, dict) else {}
    return {
        "heading": _clean(col.get("heading"), MAX_TITLE_CHARS),
        "bullets": _clean_bullets(col.get("bullets"), 5),
    }


def _clean_items(value: Any) -> list[dict[str, str]]:
    if not isinstance(value, list):
        return []
    items = []
    for item in value[:MAX_ITEMS]:
        if not isinstance(item, dict):
            continue
        heading = _clean(item.get("heading"), MAX_ITEM_HEADING_CHARS)
        text = _clean(item.get("text"), MAX_ITEM_TEXT_CHARS)
        if heading or text:
            items.append({"heading": heading, "text": text})
    return items


def _clean_image(value: Any) -> dict[str, str] | None:
    """Keep only images we cached ourselves (ids are content hashes, never paths)."""
    if not isinstance(value, dict):
        return None
    image_id = str(value.get("id") or "")
    if not IMAGE_ID_RE.match(image_id):
        return None
    return {
        "id": image_id,
        "url": image_url(image_id),
        "credit": _clean(value.get("credit"), 200),
        "source": _clean(value.get("source"), 500),
    }


def normalize_deck(raw: Any) -> dict[str, Any]:
    """Coerce model or client JSON into a safe, bounded deck structure."""
    if not isinstance(raw, dict):
        raise PresentationError("Deck must be a JSON object.")
    slides_in = raw.get("slides")
    if not isinstance(slides_in, list):
        raise PresentationError("Deck has no slides.")

    slides: list[dict[str, Any]] = []
    for item in slides_in[: MAX_SLIDES + 5]:
        if not isinstance(item, dict):
            continue
        layout = str(item.get("layout") or "bullets").strip().lower()
        if layout not in LAYOUTS:
            layout = "bullets"
        slide = {
            "layout": layout,
            "title": _clean(item.get("title"), MAX_TITLE_CHARS),
            "subtitle": _clean(item.get("subtitle"), MAX_BULLET_CHARS),
            "bullets": _clean_bullets(item.get("bullets")),
            "items": _clean_items(item.get("items"))[: ITEM_LAYOUTS.get(layout, MAX_ITEMS)],
            "left": _clean_column(item.get("left")),
            "right": _clean_column(item.get("right")),
            "image_query": _clean(item.get("image_query"), MAX_QUERY_CHARS),
            "image": _clean_image(item.get("image")),
            "notes": _clean(item.get("notes"), MAX_NOTES_CHARS),
        }
        if layout == "two_column" and not (slide["left"]["bullets"] or slide["right"]["bullets"]):
            slide["layout"] = "bullets"
        if layout in ITEM_LAYOUTS and len(slide["items"]) < 2:
            slide["layout"] = "bullets"
            if not slide["bullets"]:
                slide["bullets"] = [": ".join(filter(None, (i["heading"], i["text"]))) for i in slide["items"]]
        if not slide["title"] and not slide["bullets"]:
            continue
        slides.append(slide)

    if not slides:
        raise PresentationError("Deck has no usable slides.")
    theme = str(raw.get("theme") or "").strip().lower()
    return {
        "title": _clean(raw.get("title"), MAX_TITLE_CHARS) or slides[0]["title"],
        "subtitle": _clean(raw.get("subtitle"), MAX_BULLET_CHARS),
        "presenter": _clean(raw.get("presenter"), MAX_PRESENTER_CHARS),
        "theme": theme if theme in THEMES else DEFAULT_THEME,
        "image_query": _clean(raw.get("image_query"), MAX_QUERY_CHARS),
        "cover_image": _clean_image(raw.get("cover_image")),
        "slides": slides,
    }


def requested_slide_count(text: str) -> int | None:
    """A slide count the user wrote in the request ("8-slide deck", "10 слайдов"), if in range."""
    for match in _SLIDE_COUNT_RE.finditer(text or ""):
        count = int(match.group(1))
        if MIN_SLIDES <= count <= MAX_SLIDES:
            return count
    return None


def requested_minutes(text: str) -> float | None:
    """Talk length the user asked for ("7–10 minutes" -> 8.5), if plausible."""
    match = _DURATION_RE.search(text or "")
    if not match:
        return None
    low = int(match.group(1))
    high = int(match.group(2) or low)
    minutes = (low + high) / 2
    return minutes if 1 <= minutes <= 60 else None


def generate_deck(
    *,
    topic: str,
    details: str = "",
    slide_count: int = 8,
    tone: str = "academic",
    language: str = "auto",
    theme: str = DEFAULT_THEME,
    images: bool = True,
) -> tuple[dict[str, Any], dict[str, Any]]:
    """Ask Claude for a deck and attach topic photos. Returns (deck, diagnostics); raises PresentationError."""
    topic = _clean(topic, MAX_TOPIC_CHARS)
    if not topic:
        raise PresentationError("Topic is required.")
    details = (details or "").strip()[:MAX_DETAILS_CHARS]
    slide_count = requested_slide_count(f"{topic}\n{details}") or int(slide_count)
    slide_count = max(MIN_SLIDES, min(MAX_SLIDES, slide_count))
    tone_text = TONES.get(tone, TONES["academic"])
    language_text = LANGUAGES.get(language, LANGUAGES["auto"])

    parts = [
        f"Topic: {topic}",
        f"Write exactly {slide_count - 1} slides in \"slides\" (a title slide is added automatically).",
        f"Audience and tone: {tone_text}.",
        f"Write all text in {language_text}.",
    ]
    minutes = requested_minutes(f"{topic}\n{details}")
    if minutes:
        per_slide = round(minutes * _NOTES_WORDS_PER_MINUTE / (slide_count - 1) / 10) * 10
        parts.append(
            f"The talk lasts about {minutes:g} minutes: write about {per_slide} words of speaker notes "
            "on every slide so the whole script fills that time."
        )
    if not images:
        parts.append('Photos are turned off: do not use the "image" layout and leave every image_query empty.')
    if details:
        parts.append(
            "Use this material from the user (brief, notes, or source text) as the main basis:\n"
            f"<material>\n{details}\n</material>"
        )

    data, diag = generate_json(
        system_prompt=_SYSTEM_PROMPT,
        user_prompt="\n".join(parts),
        temperature=0.5,
        max_tokens=16000,
        timeout_s=300,
    )
    if data is None:
        raise PresentationError(diag.get("error_message") or "Presentation generation failed.")
    data["theme"] = theme
    deck = normalize_deck(data)
    diag["images_found"] = attach_images(deck, layouts=PHOTO_LAYOUTS) if images else 0
    for slide in deck["slides"]:
        if slide["layout"] == "image" and not slide["image"]:
            slide["layout"] = "bullets"
    return deck, diag


# --- PPTX rendering -------------------------------------------------------

_SLIDE_W_IN = 13.333
_SLIDE_H_IN = 7.5
_CONTENT_BOTTOM_IN = 6.85
# Average glyph width in em, measured on Cyrillic/Latin renders of these fonts.
_HEADING_EM = 0.62
_HEADING_UPPER_EM = 0.72
_BODY_EM = 0.5
_BULLET_PREFIX = "\u2022  "


def _rgb(hex_color: str):
    from pptx.dml.color import RGBColor

    return RGBColor.from_string(hex_color)


def _mix(base: str, other: str, amount: float) -> str:
    """Blend hex color `base` toward `other` (0 = base, 1 = other)."""
    a = [int(base[i : i + 2], 16) for i in (0, 2, 4)]
    b = [int(other[i : i + 2], 16) for i in (0, 2, 4)]
    return "".join(f"{round(x + (y - x) * amount):02X}" for x, y in zip(a, b))


def _wrapped_lines(text: str, chars_per_line: int) -> int:
    lines, current = 1, 0
    for word in text.split():
        size = len(word)
        if current == 0:
            current = size
        elif current + 1 + size <= chars_per_line:
            current += 1 + size
        else:
            lines += 1
            current = size
        while current > chars_per_line:
            lines += 1
            current -= chars_per_line
    return lines


def _text_height(
    lines: list[str], width_in: float, size: int, *, em: float, spacing: float, prefix: str = "", line: float = 1.2
) -> float:
    """Estimated rendered height in inches (python-pptx cannot measure text). `line` = line height in em."""
    chars_per_line = max(8, int((width_in - 0.2) * 72 / (size * em)))
    points = sum(
        _wrapped_lines(prefix + text, chars_per_line) * size * line + size * spacing for text in lines
    )
    return points / 72 + 0.1


def _fit_size(
    lines: list[str],
    width_in: float,
    height_in: float,
    *,
    max_size: int,
    min_size: int,
    em: float = _BODY_EM,
    spacing: float = 0.0,
    prefix: str = "",
    line: float = 1.2,
) -> int:
    longest_word = max((len(word) for text in lines for word in text.split()), default=0)
    for size in range(max_size, min_size - 1, -1):
        chars_per_line = max(8, int((width_in - 0.2) * 72 / (size * em)))
        if longest_word > chars_per_line:
            continue
        if _text_height(lines, width_in, size, em=em, spacing=spacing, prefix=prefix, line=line) <= height_in:
            return size
    return min_size


def _rect(slide, left_in, top_in, width_in, height_in, hex_color: str, *, rounded: bool = False):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    kind = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_color)
    shape.line.fill.background()
    if rounded:
        shape.adjustments[0] = 0.06
    return shape


# Preset shapes only: Quick Look and iOS skip custom geometry (freeforms) entirely.
def _triangle(slide, left_in, top_in, width_in, height_in, hex_color: str, corner: str):
    """Right triangle with the right angle at `corner` ("bl", "br", "tl" or "tr")."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.RIGHT_TRIANGLE, Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in)
    )
    xfrm = shape._element.spPr.xfrm
    if corner in ("br", "tr"):
        xfrm.set("flipH", "1")
    if corner in ("tl", "tr"):
        xfrm.set("flipV", "1")
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_color)
    shape.line.fill.background()
    return shape


def _slant(slide, left_in, width_in, offset_in, hex_color: str):
    """Full-height parallelogram whose top edge is shifted `offset_in` right of its bottom edge."""
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    shape = slide.shapes.add_shape(
        MSO_SHAPE.PARALLELOGRAM, Inches(left_in), 0, Inches(width_in), Inches(_SLIDE_H_IN)
    )
    shape.adjustments[0] = offset_in / min(width_in, _SLIDE_H_IN)
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_color)
    shape.line.fill.background()
    return shape


def _alpha(color_element, opacity: float) -> None:
    from pptx.oxml.xmlchemy import OxmlElement

    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(int(opacity * 100000)))
    color_element.append(alpha)


def _scrim(slide, hex_color: str) -> None:
    """Tint over a photo, light at the top and dense at the bottom where the title sits.
    Stacked translucent strips: Quick Look renders gradient fills with alpha as opaque."""
    from pptx.oxml.ns import qn

    def layer(top, opacity):
        shape = _rect(slide, 0, top, _SLIDE_W_IN, _SLIDE_H_IN - top, hex_color)
        _alpha(shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")), opacity)

    layer(0, 0.4)
    for i in range(8):
        layer(1.5 + i * 0.5, 0.09)


def _disc(slide, center_x_in, center_y_in, diameter_in, hex_color: str):
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    r = diameter_in / 2
    shape = slide.shapes.add_shape(
        MSO_SHAPE.OVAL, Inches(center_x_in - r), Inches(center_y_in - r), Inches(diameter_in), Inches(diameter_in)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(hex_color)
    shape.line.fill.background()
    return shape


def _ring(slide, center_x_in, center_y_in, diameter_in, hex_color: str, width_pt: float = 1.25):
    from pptx.util import Pt

    shape = _disc(slide, center_x_in, center_y_in, diameter_in, hex_color)
    shape.fill.background()
    shape.line.color.rgb = _rgb(hex_color)
    shape.line.width = Pt(width_pt)
    return shape


def _text(
    slide,
    left_in: float,
    top_in: float,
    width_in: float,
    height_in: float,
    lines: list[str],
    *,
    size: int,
    color: str,
    font: str,
    bold: bool = False,
    bullets: bool = False,
    marker_color: str = "",
    align=None,
    accent: tuple[str, str] | None = None,
):
    """`accent` = (title_accent mode, color) recolors part of the first line."""
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Inches, Pt

    box = slide.shapes.add_textbox(Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    frame = box.text_frame
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE
    for i, line in enumerate(lines):
        para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        if align is not None:
            para.alignment = align
        if bullets:
            para.space_after = Pt(size * 0.6)
            marker = para.add_run()
            marker.text = _BULLET_PREFIX
            marker.font.size = Pt(size)
            marker.font.color.rgb = _rgb(marker_color or color)
            marker.font.bold = True
            marker.font.name = font
        segments = accent_split(line, accent[0]) if accent and i == 0 else [(line, False)]
        for segment, highlighted in segments:
            run = para.add_run()
            run.text = segment
            run.font.size = Pt(size)
            run.font.color.rgb = _rgb(accent[1] if highlighted else color)
            run.font.bold = bold
            run.font.name = font
    return box


def _heading_em(t: Theme) -> float:
    return (_HEADING_UPPER_EM if t.heading_upper else _HEADING_EM) * t.heading_width


def _heading_text(t: Theme, text: str) -> tuple[str, float]:
    return (text.upper() if t.heading_upper else text), _heading_em(t)


def _heading(
    slide,
    t: Theme,
    left_in,
    top_in,
    width_in,
    max_height_in,
    text,
    *,
    max_size,
    min_size,
    color,
    align_bottom=False,
    accent_color: str = "",
) -> float:
    """Draw a theme heading sized to fit; returns its estimated height in inches."""
    text, em = _heading_text(t, text)
    line = t.heading_line
    size = _fit_size([text], width_in, max_height_in, max_size=max_size, min_size=min_size, em=em, line=line)
    height = min(max_height_in, _text_height([text], width_in, size, em=em, spacing=0, line=line))
    box_top = top_in + max_height_in - height if align_bottom else top_in
    accent = (t.title_accent, accent_color) if accent_color and t.title_accent else None
    _text(
        slide, left_in, box_top, width_in, height, [text],
        size=size, color=color, font=t.heading_font, bold=t.heading_bold, accent=accent,
    )
    return height


def _bullet_fit(lines, width_in, height_in, *, max_size) -> int:
    return _fit_size(lines, width_in, height_in, max_size=max_size, min_size=12, spacing=0.6, prefix=_BULLET_PREFIX)


def _bullet_block(slide, t: Theme, left_in, top_in, width_in, height_in, lines, *, max_size, color, marker_color, size=None):
    if not lines:
        return
    if size is None:
        size = _bullet_fit(lines, width_in, height_in, max_size=max_size)
    _text(
        slide,
        left_in,
        top_in,
        width_in,
        height_in,
        lines,
        size=size,
        color=color,
        font=t.body_font,
        bullets=True,
        marker_color=marker_color,
    )


def _circle_jpeg(path, background: str) -> io.BytesIO:
    """Center-crop to a circle on a solid background (Quick Look ignores picture geometry)."""
    from PIL import Image, ImageDraw

    size = 900
    with Image.open(path) as im:
        im = im.convert("RGB")
        side = min(im.size)
        left, top = (im.width - side) // 2, (im.height - side) // 2
        photo = im.crop((left, top, left + side, top + side)).resize((size, size), Image.LANCZOS)
    mask = Image.new("L", (size * 3, size * 3), 0)
    ImageDraw.Draw(mask).ellipse((0, 0, size * 3 - 1, size * 3 - 1), fill=255)
    mask = mask.resize((size, size), Image.LANCZOS)
    out = Image.new("RGB", (size, size), "#" + background)
    out.paste(photo, (0, 0), mask)
    buf = io.BytesIO()
    out.save(buf, "JPEG", quality=85)
    buf.seek(0)
    return buf


def _picture(
    slide, image: dict[str, str] | None, left_in, top_in, width_in, height_in, *, circle_bg: str = ""
) -> bool:
    """Place a cached photo cropped to fill the box (like CSS object-fit: cover), or as a circle."""
    from PIL import Image
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    path = image_path(image["id"]) if image else None
    if path is None:
        return False
    if circle_bg:
        pic = slide.shapes.add_picture(
            _circle_jpeg(path, circle_bg), Inches(left_in), Inches(top_in), Inches(width_in), Inches(width_in)
        )
        pic.auto_shape_type = MSO_SHAPE.OVAL
        return True
    with Image.open(path) as im:
        img_ratio = im.width / im.height
    pic = slide.shapes.add_picture(str(path), Inches(left_in), Inches(top_in), Inches(width_in), Inches(height_in))
    box_ratio = width_in / height_in
    if img_ratio > box_ratio:
        pic.crop_left = pic.crop_right = (1 - box_ratio / img_ratio) / 2
    else:
        pic.crop_top = pic.crop_bottom = (1 - img_ratio / box_ratio) / 2
    return True


def _has_image(image: dict[str, str] | None) -> bool:
    return bool(image and image_path(image.get("id", "")))


def _oval(slide, center_x_in, center_y_in, diameter_in, fill: str, label: str, label_color: str, font: str):
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Pt

    shape = _disc(slide, center_x_in, center_y_in, diameter_in, fill)
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    para = frame.paragraphs[0]
    para.alignment = PP_ALIGN.CENTER
    run = para.add_run()
    run.text = label
    run.font.size = Pt(int(diameter_in * 30))
    run.font.bold = True
    run.font.name = font
    run.font.color.rgb = _rgb(label_color)


def _clip(text: str, width_in: float, size: int) -> str:
    limit = max(8, int((width_in - 0.2) * 72 / (size * 0.62)))
    return text if len(text) <= limit else text[: limit - 1].rstrip() + "…"


def _chrome(slide, t: Theme, number: int, total: int, deck_title: str, right_in: float = _SLIDE_W_IN) -> None:
    """Page number plus the design's running header/footer on a content slide."""
    from pptx.enum.text import PP_ALIGN

    design = t.design
    rule = _mix(t.bg, t.ink, 0.22)
    if design == "magazine":
        label_w = right_in - 0.8 - 1.9
        _text(
            slide, 0.8, 0.2, label_w, 0.36, [_clip(deck_title.upper(), label_w, 10)],
            size=10, color=t.muted, font=t.body_font, bold=True,
        )
        _text(
            slide, right_in - 1.6, 0.2, 1.0, 0.36, [f"{number:02d}"],
            size=10, color=t.accent, font=t.body_font, bold=True, align=PP_ALIGN.RIGHT,
        )
        _rect(slide, 0.8, 0.6, right_in - 1.4, 0.012, rule)
        return
    if design == "editorial":
        label_w = right_in - 0.8 - 2.2
        _rect(slide, 0.8, 6.98, right_in - 1.4, 0.012, rule)
        _text(
            slide, 0.8, 7.0, label_w, 0.36, [_clip(deck_title, label_w, 10)],
            size=10, color=t.muted, font=t.body_font,
        )
        _text(
            slide, right_in - 1.8, 7.0, 1.2, 0.36, [f"{number:02d} / {total:02d}"],
            size=10, color=t.accent, font=t.body_font, bold=True, align=PP_ALIGN.RIGHT,
        )
        return
    if design == "diagonal":
        if right_in >= _SLIDE_W_IN:
            _triangle(slide, 12.15, 0, _SLIDE_W_IN - 12.15, 1.15, t.accent, "tr")
        _rect(slide, 0.8, 7.13, 0.35, 0.05, t.accent)
        _text(slide, 1.2, 6.95, 1.2, 0.4, [f"{number:02d}"], size=11, color=t.muted, font=t.body_font, bold=True)
        return
    if design == "lineart":
        if right_in >= _SLIDE_W_IN:
            ring = _mix(t.bg, t.accent, 0.4)
            _ring(slide, _SLIDE_W_IN, _SLIDE_H_IN, 1.4, ring, 1)
            _ring(slide, _SLIDE_W_IN, _SLIDE_H_IN, 2.0, ring, 0.75)
        _text(slide, 0.8, 6.95, 1.2, 0.4, [f"{number:02d}"], size=11, color=t.accent, font=t.body_font, bold=True)
        return
    _text(
        slide, right_in - 1.6, _CONTENT_BOTTOM_IN + 0.1, 1.2, 0.4, [f"{number} / {total}"],
        size=11, color=t.muted, font=t.body_font, align=PP_ALIGN.RIGHT,
    )


def _blank(prs, fill: str):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(slide, 0, 0, _SLIDE_W_IN, _SLIDE_H_IN, fill)
    return slide


def _cover_frame(slide, t: Theme) -> float:
    """Decorate a classic cover/closing slide; returns the left text inset in inches."""
    if t.cover_bg.upper() == t.bg.upper():
        _rect(slide, 0, 0, 0.32, _SLIDE_H_IN, t.cover_accent)
        return 1.1
    _rect(slide, 0, _SLIDE_H_IN - 0.3, _SLIDE_W_IN, 0.3, t.cover_accent)
    return 0.9


# Photos bleed off the right edge from this x (inches).
_PHOTO_X_IN = 7.6
_COVER_PHOTO_X_IN = 7.0
# Inset photo (left, top, width, height) for designs that frame photos instead of bleeding them.
_INSET_PHOTO = (7.75, 0.6, 4.98, 6.3)


def _photo_box(t: Theme) -> tuple[float, float, float, float]:
    if t.design in ("editorial", "lineart"):
        return _INSET_PHOTO
    return (_PHOTO_X_IN, 0.0, _SLIDE_W_IN - _PHOTO_X_IN, _SLIDE_H_IN)


def _subtitle_and_byline(
    slide, t: Theme, left, top, width, subtitle, presenter, *, box_h=1.5, byline_color="", upper_byline=False
) -> float:
    """Cover subtitle with the presenter line under it; returns the bottom y."""
    bottom = top
    if subtitle:
        box_h = box_h - 0.45 if presenter else box_h
        size = _fit_size([subtitle], width, box_h, max_size=22, min_size=14)
        sub_h = min(box_h, _text_height([subtitle], width, size, em=_BODY_EM, spacing=0))
        _text(slide, left, top, width, sub_h, [subtitle], size=size, color=t.cover_muted, font=t.body_font)
        bottom = top + sub_h + 0.15
    if presenter:
        text = presenter.upper() if upper_byline else presenter
        size = _fit_size([text], width, 0.5, max_size=16 if not upper_byline else 13, min_size=11)
        _text(
            slide, left, bottom, width, 0.5, [text],
            size=size, color=byline_color or t.cover_ink, font=t.body_font, bold=True,
        )
        bottom += 0.5
    return bottom


def _cover_classic(prs, t: Theme, title, subtitle, image, presenter) -> Any:
    slide = _blank(prs, t.cover_bg)
    left = _cover_frame(slide, t)
    has_photo = _picture(slide, image, _COVER_PHOTO_X_IN, 0, _SLIDE_W_IN - _COVER_PHOTO_X_IN, _SLIDE_H_IN)
    width = (_COVER_PHOTO_X_IN - 0.5 if has_photo else _SLIDE_W_IN - 1.0) - left
    _heading(
        slide, t, left, 1.0, width, 3.5, title,
        max_size=40 if has_photo else 44, min_size=24, color=t.cover_ink, align_bottom=True,
        accent_color=t.cover_accent,
    )
    _rect(slide, left + 0.1, 4.68, 1.4, 0.07, t.cover_accent)
    _subtitle_and_byline(slide, t, left, 4.95, width, subtitle, presenter)
    return slide


def _cover_magazine(prs, t: Theme, title, subtitle, image, presenter) -> Any:
    """Running header strip, oversized title, framed photo on the right."""
    slide = _blank(prs, t.cover_bg)
    rule = _mix(t.cover_bg, t.cover_ink, 0.3)
    has_photo = _picture(slide, image, 7.45, 1.05, 5.28, 5.6)
    if presenter:
        _text(
            slide, 0.6, 0.27, 8.0, 0.4, [_clip(presenter.upper(), 8.0, 11)],
            size=11, color=t.cover_muted, font=t.body_font, bold=True,
        )
    _rect(slide, 0.6, 0.75, 12.13, 0.015, rule)
    _rect(slide, 0.6, 6.95, 12.13, 0.015, rule)
    width = 6.45 if has_photo else 11.5
    _heading(
        slide, t, 0.6, 1.05, width, 3.75, title,
        max_size=48 if has_photo else 54, min_size=24, color=t.cover_ink, align_bottom=True,
        accent_color=t.cover_accent,
    )
    _subtitle_and_byline(slide, t, 0.6, 5.0, width, subtitle, "", box_h=1.7)
    return slide


_DIAGONAL_RUN_IN = 2.2


def _diagonal_band(slide, t: Theme, top_x: float, *, fill_right: bool) -> None:
    """Accent stripe running from (top_x, 0) down-left to (top_x - 2.2, 7.5)."""
    bottom_x = top_x - _DIAGONAL_RUN_IN
    if fill_right:
        tint = _mix(t.cover_bg, t.cover_accent, 0.14)
        _rect(slide, top_x + 0.3, 0, _SLIDE_W_IN - top_x - 0.3, _SLIDE_H_IN, tint)
        _triangle(slide, bottom_x + 0.3, 0, _DIAGONAL_RUN_IN + 0.01, _SLIDE_H_IN, tint, "br")
    _slant(slide, bottom_x, _DIAGONAL_RUN_IN + 0.35, _DIAGONAL_RUN_IN, t.cover_accent)


def _cover_diagonal(prs, t: Theme, title, subtitle, image, presenter) -> Any:
    """Photo cut by a slanted accent stripe."""
    slide = _blank(prs, t.cover_bg)
    has_photo = _picture(slide, image, 5.8, 0, _SLIDE_W_IN - 5.8, _SLIDE_H_IN)
    if has_photo:
        _triangle(slide, 5.79, 0, _DIAGONAL_RUN_IN + 0.02, _SLIDE_H_IN, t.cover_bg, "tl")
    _diagonal_band(slide, t, 8.0, fill_right=not has_photo)
    _rect(slide, 0.8, 1.2, 0.9, 0.08, t.cover_accent)
    width = 5.4
    title_h = _heading(
        slide, t, 0.8, 1.5, width, 3.3, title, max_size=40, min_size=22, color=t.cover_ink,
        accent_color=t.cover_accent,
    )
    _subtitle_and_byline(slide, t, 0.8, 1.5 + title_h + 0.3, width, subtitle, presenter, byline_color=t.cover_accent)
    return slide


def _cover_editorial(prs, t: Theme, title, subtitle, image, presenter) -> Any:
    """Serif title, soft circle, photo with an offset accent block, meta row at the bottom."""
    slide = _blank(prs, t.cover_bg)
    _disc(slide, 0.9, 0.5, 4.4, _mix(t.cover_bg, t.cover_accent, 0.1))
    has_photo = _has_image(image)
    if has_photo:
        _rect(slide, 7.85, 1.0, 4.9, 5.75, t.cover_accent)
        _picture(slide, image, 7.55, 0.7, 4.9, 5.75)
    width = 6.2 if has_photo else 11.3
    _rect(slide, 0.95, 1.3, 0.6, 0.045, t.cover_accent)
    _heading(
        slide, t, 0.9, 1.5, width, 3.4, title, max_size=46, min_size=24, color=t.cover_ink, align_bottom=True,
        accent_color=t.cover_accent,
    )
    _subtitle_and_byline(slide, t, 0.9, 5.1, width, subtitle, "", box_h=1.2)
    if presenter:
        _rect(slide, 0.9, 6.55, min(width, 6.0), 0.012, _mix(t.cover_bg, t.cover_ink, 0.3))
        _text(
            slide, 0.9, 6.62, width, 0.4, [_clip(presenter, width, 13)],
            size=13, color=t.cover_muted, font=t.body_font,
        )
    return slide


def _cover_fullbleed(prs, t: Theme, title, subtitle, image, presenter) -> Any:
    """Full-slide photo under a tinted overlay, large title bottom-left."""
    slide = _blank(prs, t.cover_bg)
    if _picture(slide, image, 0, 0, _SLIDE_W_IN, _SLIDE_H_IN):
        _scrim(slide, t.cover_bg)
    else:
        _disc(slide, 11.6, 1.6, 6.5, _mix(t.cover_bg, t.cover_accent, 0.18))
    width = 10.5
    title_h = _heading(
        slide, t, 0.8, 1.5, width, 3.7, title, max_size=54, min_size=26, color=t.cover_ink, align_bottom=True,
        accent_color=t.cover_accent,
    )
    _rect(slide, 0.88, 5.2 - title_h - 0.22, 1.1, 0.09, t.cover_accent)
    _subtitle_and_byline(slide, t, 0.8, 5.4, 9.5, subtitle, presenter, box_h=1.6, upper_byline=True)
    return slide


def _cover_lineart(prs, t: Theme, title, subtitle, image, presenter) -> Any:
    """Thin rings; the photo sits in a circle."""
    slide = _blank(prs, t.cover_bg)
    ring = t.cover_accent
    has_photo = _has_image(image)
    if has_photo:
        _picture(slide, image, 7.4, 1.25, 5.0, 5.0, circle_bg=t.cover_bg)
        _ring(slide, 10.35, 3.45, 6.7, _mix(t.cover_bg, ring, 0.35), 0.75)
        _ring(slide, 9.9, 3.75, 5.75, ring, 1.25)
        _disc(slide, 7.87, 1.72, 0.32, ring)
    else:
        for i, d in enumerate((3.0, 4.8, 6.6, 8.4)):
            _ring(slide, 12.6, 3.75, d, _mix(t.cover_bg, ring, 0.75 - i * 0.15), 1)
    width = 6.0 if has_photo else 8.0
    _heading(
        slide, t, 0.9, 1.2, width, 3.4, title, max_size=42, min_size=24, color=t.cover_ink, align_bottom=True,
        accent_color=t.cover_accent,
    )
    _rect(slide, 0.95, 4.78, 0.7, 0.04, ring)
    _subtitle_and_byline(slide, t, 0.9, 5.0, width, subtitle, presenter, byline_color=ring)
    return slide


_COVERS = {
    "classic": _cover_classic,
    "magazine": _cover_magazine,
    "diagonal": _cover_diagonal,
    "editorial": _cover_editorial,
    "fullbleed": _cover_fullbleed,
    "lineart": _cover_lineart,
}


def _render_cover(
    prs, t: Theme, title: str, subtitle: str, image: dict[str, str] | None = None, presenter: str = ""
) -> Any:
    return _COVERS.get(t.design, _cover_classic)(prs, t, title, subtitle, image, presenter)


def _render_title_block(slide, t: Theme, title: str, width_in: float = 11.7) -> float:
    """Draw the slide title in the theme's design; returns where body content may start."""
    design = t.design

    def heading(left, top, width, color):
        return _heading(
            slide, t, left, top, width, 1.6, title, max_size=30 if design != "classic" else 32, min_size=20,
            color=color, accent_color=t.accent,
        )

    if design == "magazine":
        return 0.9 + heading(0.8, 0.9, width_in, t.ink) + 0.3
    if design == "fullbleed":
        _rect(slide, 0.88, 0.5, 0.9, 0.08, t.accent)
        return 0.72 + heading(0.8, 0.72, width_in, t.ink) + 0.3
    if design == "editorial":
        return 0.6 + heading(0.8, 0.6, width_in, t.ink) + 0.35
    if design == "diagonal":
        line_top = 0.5 + heading(0.8, 0.5, width_in, t.ink) + 0.1
        _rect(slide, 0.8, line_top + 0.025, width_in, 0.012, _mix(t.bg, t.ink, 0.22))
        _rect(slide, 0.8, line_top, 1.3, 0.06, t.accent)
        return line_top + 0.4
    if design == "lineart":
        title_h = heading(1.15, 0.5, width_in - 0.35, t.col_heading)
        _rect(slide, 0.85, 0.6, 0.06, max(0.3, title_h - 0.2), t.accent)
        line_top = 0.5 + title_h + 0.12
        _rect(slide, 0.8, line_top, width_in, 0.012, _mix(t.bg, t.ink, 0.15))
        return line_top + 0.35
    _rect(slide, 0, 0, 0.18, _SLIDE_H_IN, t.accent)
    rule_top = 0.45 + _heading(
        slide, t, 0.8, 0.45, width_in, 1.6, title, max_size=32, min_size=22, color=t.ink, accent_color=t.accent
    ) + 0.08
    _rect(slide, 0.9, rule_top, 1.2, 0.06, t.accent)
    return rule_top + 0.35


def _render_bullets(prs, t: Theme, data: dict[str, Any]) -> Any:
    slide = _blank(prs, t.bg)
    top = _render_title_block(slide, t, data["title"])
    lines = data["bullets"] or ([data["subtitle"]] if data["subtitle"] else [])
    _bullet_block(
        slide, t, 0.8, top, 11.7, _CONTENT_BOTTOM_IN - top, lines, max_size=24, color=t.ink, marker_color=t.accent
    )
    return slide


def _render_image(prs, t: Theme, data: dict[str, Any]) -> Any:
    if not _has_image(data.get("image")):
        return _render_bullets(prs, t, data)
    slide = _blank(prs, t.bg)
    left, top_in, width_in, height_in = _photo_box(t)
    if t.design == "editorial":
        _rect(slide, left + 0.2, top_in + 0.2, width_in, height_in, t.accent)
    _picture(slide, data["image"], left, top_in, width_in, height_in)
    width = left - 0.8 - 0.5
    top = _render_title_block(slide, t, data["title"], width_in=width)
    lines = data["bullets"] or ([data["subtitle"]] if data["subtitle"] else [])
    _bullet_block(
        slide, t, 0.8, top, width, _CONTENT_BOTTOM_IN - top, lines, max_size=22, color=t.ink, marker_color=t.accent
    )
    return slide


def _uniform_size(
    texts: list[str], width_in: float, height_in: float, *, max_size: int, min_size: int, em: float, line: float = 1.2
) -> int:
    return min(
        (
            _fit_size([s], width_in, height_in, max_size=max_size, min_size=min_size, em=em, line=line)
            for s in texts
            if s
        ),
        default=max_size,
    )


def _render_cards(prs, t: Theme, data: dict[str, Any]) -> Any:
    """2-3 cards in a row (number on top), or 4 cards in a 2x2 grid (number on the left)."""
    slide = _blank(prs, t.bg)
    items = data["items"]
    cols = 2 if len(items) == 4 else len(items)
    rows = (len(items) + cols - 1) // cols
    grid = rows > 1
    top = _render_title_block(slide, t, data["title"]) + (0.0 if grid else 0.1)
    gap = 0.25 if grid else 0.3
    card_w = (11.7 - gap * (cols - 1)) / cols
    card_h = (_CONTENT_BOTTOM_IN - 0.15 - top - gap * (rows - 1)) / rows
    # Offsets inside a card: text column x, heading y, bottom padding, accent rule y.
    pad_x = 1.1 if grid else 0.3
    head_y = 0.22 if grid else 1.15
    pad_bottom = 0.12 if grid else 0.2
    rule_y = 0.22 if grid else 0.3
    inner_w = card_w - pad_x - 0.3
    head_texts = [_heading_text(t, i["heading"])[0] for i in items]
    em = _heading_em(t)
    head_size = _uniform_size(head_texts, inner_w, 1.1, max_size=22, min_size=14, em=em, line=t.heading_line)
    if grid:
        # Grid cards are short: prefer one-line headings, but not at the cost of tiny type.
        one_line = _uniform_size(head_texts, inner_w, 0.5, max_size=20, min_size=18, em=em, line=t.heading_line)
        fits = all(_text_height([h], inner_w, one_line, em=em, spacing=0, line=t.heading_line) <= 0.5 for h in head_texts if h)
        head_size = one_line if fits else _uniform_size(head_texts, inner_w, 0.9, max_size=20, min_size=14, em=em, line=t.heading_line)
    head_h = max((_text_height([h], inner_w, head_size, em=em, spacing=0, line=t.heading_line) for h in head_texts if h), default=0.0)
    text_y = head_y + head_h + 0.05
    text_h = card_h - pad_bottom - text_y
    text_size = _uniform_size([i["text"] for i in items], inner_w, text_h, max_size=20, min_size=11, em=_BODY_EM)
    text_size = min(text_size, head_size)

    for idx, item in enumerate(items):
        left = 0.8 + (idx % cols) * (card_w + gap)
        card_top = top + (idx // cols) * (card_h + gap)
        _rect(slide, left, card_top, card_w, card_h, t.pale, rounded=True)
        _rect(slide, left + 0.3, card_top + rule_y, 0.5, 0.06, t.accent)
        _text(
            slide, left + 0.3, card_top + rule_y + 0.12, 0.9, 0.6, [f"{idx + 1:02d}"],
            size=26 if grid else 24, color=t.accent, font=t.heading_font, bold=True,
        )
        if item["heading"]:
            _text(
                slide, left + pad_x, card_top + head_y, inner_w, head_h, [head_texts[idx]],
                size=head_size, color=t.col_heading, font=t.heading_font, bold=True,
            )
        if item["text"]:
            _text(
                slide, left + pad_x, card_top + text_y, inner_w, text_h, [item["text"]],
                size=text_size, color=t.ink, font=t.body_font,
            )
    return slide


def _render_timeline(prs, t: Theme, data: dict[str, Any]) -> Any:
    from pptx.enum.text import PP_ALIGN

    slide = _blank(prs, t.bg)
    top = _render_title_block(slide, t, data["title"])
    items = data["items"]
    col_w = 11.7 / len(items)
    dot = 0.62
    inner_w = col_w - 0.3
    head_texts = [_heading_text(t, i["heading"])[0] for i in items]
    em = _heading_em(t)
    head_size = _uniform_size(head_texts, inner_w, 1.0, max_size=22, min_size=13, em=em, line=t.heading_line)
    head_h = max((_text_height([h], inner_w, head_size, em=em, spacing=0, line=t.heading_line) for h in head_texts if h), default=0.0)
    # Body below the dots: heading, gap, text. Size the text for the room left, then center the block.
    fixed_h = dot + 0.25 + head_h + 0.1
    room = _CONTENT_BOTTOM_IN - 0.1 - top
    texts = [i["text"] for i in items]
    text_size = _uniform_size(texts, inner_w, room - fixed_h, max_size=19, min_size=11, em=_BODY_EM)
    text_h = max((_text_height([s], inner_w, text_size, em=_BODY_EM, spacing=0) for s in texts if s), default=0.0)
    block_top = top + max(0.0, (room - fixed_h - text_h) / 2 - 0.2)
    line_y = block_top + dot / 2
    head_top = line_y + dot / 2 + 0.25
    text_top = head_top + head_h + 0.1
    _rect(slide, 0.8 + col_w / 2, line_y - 0.025, col_w * (len(items) - 1), 0.05, t.accent)

    for idx, item in enumerate(items):
        center = 0.8 + col_w * idx + col_w / 2
        _oval(slide, center, line_y, dot, t.accent, str(idx + 1), t.bg, t.heading_font)
        left = center - inner_w / 2
        if item["heading"]:
            _text(
                slide, left, head_top, inner_w, head_h, [head_texts[idx]],
                size=head_size, color=t.col_heading, font=t.heading_font, bold=True, align=PP_ALIGN.CENTER,
            )
        if item["text"]:
            _text(
                slide, left, text_top, inner_w, text_h, [item["text"]],
                size=text_size, color=t.ink, font=t.body_font, align=PP_ALIGN.CENTER,
            )
    return slide


def _render_two_column(prs, t: Theme, data: dict[str, Any]) -> Any:
    slide = _blank(prs, t.bg)
    top = _render_title_block(slide, t, data["title"])
    columns = (data["left"], data["right"])
    inner_w = 5.0
    headings = [_heading_text(t, c["heading"])[0] for c in columns if c["heading"]]
    em = _heading_em(t)
    heading_size = min(
        (_fit_size([h], inner_w, 1.0, max_size=22, min_size=16, em=em, line=t.heading_line) for h in headings),
        default=22,
    )
    heading_h = max((_text_height([h], inner_w, heading_size, em=em, spacing=0, line=t.heading_line) for h in headings), default=0.0)
    body_top = top + 0.2 + (heading_h + 0.1 if headings else 0)
    body_h = _CONTENT_BOTTOM_IN - 0.2 - body_top
    body_size = min(_bullet_fit(c["bullets"], inner_w, body_h, max_size=20) for c in columns)

    for idx, col in enumerate(columns):
        left = 0.8 + idx * 6.0
        _rect(slide, left, top, 5.7, _CONTENT_BOTTOM_IN - top, t.pale, rounded=True)
        if col["heading"]:
            _text(
                slide,
                left + 0.35,
                top + 0.2,
                inner_w,
                heading_h,
                [_heading_text(t, col["heading"])[0]],
                size=heading_size,
                color=t.col_heading,
                font=t.heading_font,
                bold=True,
            )
        _bullet_block(
            slide,
            t,
            left + 0.35,
            body_top,
            inner_w,
            body_h,
            col["bullets"],
            max_size=20,
            size=body_size,
            color=t.ink,
            marker_color=t.accent,
        )
    return slide


def _render_section(prs, t: Theme, data: dict[str, Any]) -> Any:
    if t.design != "classic":
        return _COVERS[t.design](prs, t, data["title"], data["subtitle"], data.get("image"), "")
    slide = _blank(prs, t.pale)
    has_photo = _picture(slide, data.get("image"), _COVER_PHOTO_X_IN, 0, _SLIDE_W_IN - _COVER_PHOTO_X_IN, _SLIDE_H_IN)
    width = _COVER_PHOTO_X_IN - 1.3 - 0.5 if has_photo else 11
    title_h = _heading(
        slide, t, 1.3, 1.4, width, 2.9, data["title"], max_size=40, min_size=26, color=t.ink, align_bottom=True
    )
    bar_bottom = 4.3
    if data["subtitle"]:
        size = _fit_size([data["subtitle"]], width, 1.4, max_size=20, min_size=14)
        _text(slide, 1.3, 4.4, width, 1.4, [data["subtitle"]], size=size, color=t.muted, font=t.body_font)
        bar_bottom = 4.4 + _text_height([data["subtitle"]], width, size, em=_BODY_EM, spacing=0)
    _rect(slide, 0.9, 4.3 - title_h, 0.12, bar_bottom - (4.3 - title_h), t.accent)
    return slide


def _closing_frame(slide, t: Theme) -> tuple[float, float]:
    """Decorate the closing slide in the theme's design; returns (left, width) for the text."""
    design = t.design
    if design == "magazine":
        rule = _mix(t.cover_bg, t.cover_ink, 0.3)
        _rect(slide, 0.6, 0.6, 12.13, 0.015, rule)
        _rect(slide, 0.6, 6.95, 12.13, 0.015, rule)
        return 0.6, 12.1
    if design == "diagonal":
        _diagonal_band(slide, t, 10.6, fill_right=True)
        return 0.8, 7.6
    if design == "editorial":
        _disc(slide, 0.9, 0.5, 4.4, _mix(t.cover_bg, t.cover_accent, 0.1))
        return 0.9, 11.3
    if design == "fullbleed":
        _disc(slide, 11.6, 1.6, 6.5, _mix(t.cover_bg, t.cover_accent, 0.18))
        return 0.8, 10.0
    if design == "lineart":
        for i, d in enumerate((3.0, 4.8, 6.6, 8.4)):
            _ring(slide, 12.6, 3.75, d, _mix(t.cover_bg, t.cover_accent, 0.75 - i * 0.15), 1)
        return 0.9, 7.6
    left = _cover_frame(slide, t)
    return left, _SLIDE_W_IN - left - 1.0


def _render_closing(prs, t: Theme, data: dict[str, Any]) -> Any:
    slide = _blank(prs, t.cover_bg)
    left, width = _closing_frame(slide, t)
    title_h = _heading(
        slide, t, left, 0.9, width, 1.8, data["title"], max_size=38, min_size=26, color=t.cover_ink,
        accent_color=t.cover_accent,
    )
    top = 0.9 + title_h + 0.35
    lines = data["bullets"] or ([data["subtitle"]] if data["subtitle"] else [])
    _bullet_block(
        slide,
        t,
        left,
        top,
        width,
        _CONTENT_BOTTOM_IN - top,
        lines,
        max_size=24,
        color=t.cover_ink,
        marker_color=t.cover_accent,
    )
    return slide


_RENDERERS = {
    "bullets": _render_bullets,
    "image": _render_image,
    "cards": _render_cards,
    "timeline": _render_timeline,
    "two_column": _render_two_column,
    "section": _render_section,
    "closing": _render_closing,
}
_CHROME_LAYOUTS = ("bullets", "image", "two_column", "cards", "timeline")


def _apple_compat(prs) -> None:
    """python-pptx adds the notes master part but never lists it in presentation.xml, and keeps
    the template's 4:3 size type. PowerPoint tolerates both; Keynote and iOS Quick Look refuse the file."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement

    root = prs.part._element
    size = root.find(qn("p:sldSz"))
    if size is not None:
        size.attrib.pop("type", None)
    if root.find(qn("p:notesMasterIdLst")) is not None:
        return
    rid = next((rel.rId for rel in prs.part.rels.values() if rel.reltype == RT.NOTES_MASTER), None)
    if rid is None:
        return
    entry = OxmlElement("p:notesMasterId")
    entry.set(qn("r:id"), rid)
    id_list = OxmlElement("p:notesMasterIdLst")
    id_list.append(entry)
    root.find(qn("p:sldMasterIdLst")).addnext(id_list)


def _notes_with_credit(notes: str, image: dict[str, str] | None) -> str:
    if not image or not image.get("credit"):
        return notes
    credit = image["credit"] + (f" — {image['source']}" if image.get("source") else "")
    return f"{notes}\n\n{credit}".strip()


def build_pptx(deck: dict[str, Any]) -> bytes:
    """Render a normalized deck to .pptx bytes (16:9) in the deck's theme."""
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:  # pragma: no cover - dependency guard
        raise PresentationError("python-pptx is not installed on the server.") from exc

    deck = normalize_deck(deck)
    theme = get_theme(deck["theme"])
    prs = Presentation()
    prs.slide_width = Inches(_SLIDE_W_IN)
    prs.slide_height = Inches(_SLIDE_H_IN)

    total = len(deck["slides"]) + 1
    cover = _render_cover(prs, theme, deck["title"], deck["subtitle"], deck["cover_image"], deck["presenter"])
    cover_notes = _notes_with_credit("", deck["cover_image"])
    if cover_notes:
        cover.notes_slide.notes_text_frame.text = cover_notes
    for number, data in enumerate(deck["slides"], start=2):
        slide = _RENDERERS[data["layout"]](prs, theme, data)
        if data["layout"] in _CHROME_LAYOUTS:
            right = _SLIDE_W_IN
            if data["layout"] == "image" and _has_image(data["image"]):
                right = _photo_box(theme)[0] - 0.2
            _chrome(slide, theme, number, total, deck["title"], right_in=right)
        notes = _notes_with_credit(data["notes"], data["image"] if data["layout"] in PHOTO_LAYOUTS else None)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    _apple_compat(prs)
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def deck_filename(deck: dict[str, Any]) -> str:
    title = str(deck.get("title") or "presentation")
    slug = re.sub(r"[^\w\s-]", "", title, flags=re.UNICODE).strip()
    slug = re.sub(r"[\s_-]+", "-", slug)[:60].strip("-")
    return f"{slug or 'presentation'}.pptx"
