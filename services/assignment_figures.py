"""Assignment figures: uploads, Commons auto-fetch, [[FIGURE:...]] DOCX embedding."""

from __future__ import annotations

import io
import re
import uuid
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

import requests
from docx import Document
from docx.shared import Inches

IMAGE_EXTENSIONS = frozenset({".png", ".jpg", ".jpeg", ".webp", ".gif"})

# [[FIGURE:filename.jpg|Optional caption]]
FIGURE_MARKER_RE = re.compile(
    r"\[\[\s*FIGURE\s*:\s*(?P<name>[^\]|]+?)(?:\s*\|\s*(?P<caption>[^\]]*?))?\s*\]\]",
    re.IGNORECASE,
)

_NOTE_IMAGE_SIGNAL_RE = re.compile(
    r"\b(photo|photos|image|images|picture|pictures|figure|figures|фото|картинк)\b",
    re.IGNORECASE,
)

_BRIEF_PHOTOS_RE = re.compile(
    r"\binclude\s+(?:a\s+)?(?:photos?|images?|pictures?|figures?|map)\b"
    r"|\bphotos?\s+that\s+illustrate\b"
    r"|\binsert\s+(?:photos?|images?|diagrams?|figures?)\b",
    re.IGNORECASE,
)

_INSERT_PHOTOS_RE = re.compile(
    r"\[\s*Insert\s+(?P<kind>photos?|images?|diagrams?|figures?)\s*:\s*(?P<body>[^\]]+)\]",
    re.IGNORECASE,
)

_HEADING_LINE_RE = re.compile(r"^##\s+(.+)$", re.MULTILINE)

_MAX_FIGURE_WIDTH_IN = 5.5
_MAX_IMAGES_PER_SECTION = 2
_COMMONS_UA = "DocmaxxingAssignmentBot/1.0 (academic formatter; local)"
_COMMONS_API = "https://commons.wikimedia.org/w/api.php"


def is_image_filename(name: str) -> bool:
    return Path(str(name or "")).suffix.lower() in IMAGE_EXTENSIONS


def note_asks_for_images(note: str | None) -> bool:
    return bool(_NOTE_IMAGE_SIGNAL_RE.search(str(note or "")))


def brief_requires_photos(*texts: str | None) -> bool:
    blob = "\n".join(str(t or "") for t in texts)
    if not blob.strip():
        return False
    if _INSERT_PHOTOS_RE.search(blob):
        return True
    return bool(_BRIEF_PHOTOS_RE.search(blob))


def collect_assignment_images(files: list[Any]) -> list[dict[str, str]]:
    """Return image attachment metadata from project files."""
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for file in files or []:
        original = str(getattr(file, "original_filename", None) or getattr(file, "filename", "") or "")
        storage = str(getattr(file, "storage_path", "") or "")
        file_id = str(getattr(file, "id", "") or "")
        if not is_image_filename(original) and not is_image_filename(storage):
            continue
        path = Path(storage)
        if not path.is_file():
            continue
        key = str(path.resolve())
        if key in seen:
            continue
        seen.add(key)
        out.append(
            {
                "file_id": file_id,
                "filename": Path(original).name or path.name,
                "storage_path": str(path),
                "source": "upload",
            }
        )
    return out


def _figure_lookup_keys(name: str) -> list[str]:
    """Keys for matching FIGURE filenames after markdown may strip ``_`` / ``-``."""
    raw = Path(str(name or "")).name
    if not raw:
        return []
    stem = Path(raw).stem
    ext = Path(raw).suffix.lower()
    keys = {
        raw.lower(),
        f"{stem.lower()}{ext}",
        f"{re.sub(r'[^a-z0-9]', '', stem.lower())}{ext}",
    }
    return [k for k in keys if k]


def images_by_filename(images: list[dict[str, Any]] | None) -> dict[str, Path]:
    mapping: dict[str, Path] = {}
    for item in images or []:
        if not isinstance(item, dict):
            continue
        path = Path(str(item.get("storage_path") or item.get("path") or ""))
        if not path.is_file():
            continue
        name = Path(str(item.get("filename") or path.name)).name
        for key in _figure_lookup_keys(name) + _figure_lookup_keys(path.name):
            mapping[key] = path
    return mapping


def protect_figure_markers(text: str) -> tuple[str, list[str]]:
    """Replace FIGURE markers with placeholders so humanizers do not rewrite them."""
    figures: list[str] = []

    def _replace(match: re.Match[str]) -> str:
        figures.append(match.group(0))
        return f"\n\n[[[FIGURE_MARK_{len(figures) - 1}]]]\n\n"

    protected = FIGURE_MARKER_RE.sub(_replace, text or "")
    protected = re.sub(r"\n{3,}", "\n\n", protected).strip()
    return protected, figures


def restore_figure_markers(text: str, figures: list[str]) -> str:
    restored = text or ""
    for index, marker in enumerate(figures):
        token = f"[[[FIGURE_MARK_{index}]]]"
        block = f"\n\n{marker}\n\n"
        if token in restored:
            restored = restored.replace(token, block)
            continue
        loose = re.compile(rf"\[+\s*\[*\s*FIGURE[_\s-]*MARK[_\s-]*{index}\s*\]*\s*\]+", re.I)
        restored = loose.sub(block, restored, count=1)
    return re.sub(r"\n{3,}", "\n\n", restored).strip()


def _docx_safe_image_stream(path: Path) -> io.BytesIO:
    """Re-encode images python-docx rejects (e.g. JPEG with ICC APP2 segments)."""
    from PIL import Image

    with Image.open(path) as im:
        im = im.convert("RGB") if im.mode not in {"RGB", "L"} else im
        buf = io.BytesIO()
        im.save(buf, format="JPEG", quality=90, optimize=True)
        buf.seek(0)
        return buf


def embed_figures_in_docx_bytes(
    docx_bytes: bytes,
    images: list[dict[str, Any]] | None,
    *,
    max_width_in: float = _MAX_FIGURE_WIDTH_IN,
) -> bytes:
    """Replace [[FIGURE:...]] paragraph text with embedded pictures."""
    mapping = images_by_filename(images)
    if not mapping or not docx_bytes:
        return docx_bytes

    doc = Document(io.BytesIO(docx_bytes))
    changed = False
    for paragraph in list(doc.paragraphs):
        raw = paragraph.text or ""
        if "FIGURE" not in raw.upper():
            continue
        match = FIGURE_MARKER_RE.search(raw)
        if not match:
            continue
        name = Path(match.group("name").strip()).name
        caption = (match.group("caption") or "").strip()
        path = None
        for key in _figure_lookup_keys(name):
            path = mapping.get(key)
            if path is not None:
                break
        if path is None:
            stem = re.sub(r"[^a-z0-9]", "", Path(name).stem.lower())
            for key, candidate in mapping.items():
                key_stem = re.sub(r"[^a-z0-9]", "", Path(key).stem.lower())
                if key_stem == stem or (stem and stem in key_stem):
                    path = candidate
                    break
        if path is None:
            continue

        for run in paragraph.runs:
            run.text = ""
        before = raw[: match.start()].strip()
        after = raw[match.end() :].strip()
        if before:
            paragraph.add_run(before + "\n")
        run = paragraph.add_run()
        try:
            try:
                run.add_picture(str(path), width=Inches(max_width_in))
            except Exception:  # noqa: BLE001
                run.add_picture(_docx_safe_image_stream(path), width=Inches(max_width_in))
            changed = True
        except Exception:  # noqa: BLE001
            paragraph.add_run(match.group(0))
            continue
        if caption:
            paragraph.add_run("\n" + caption)
        if after:
            paragraph.add_run("\n" + after)

    if not changed:
        return docx_bytes
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _section_short(title: str) -> str:
    text = str(title or "").split(":", 1)[0].strip()
    text = re.sub(r"^\d+[\.\)]\s*", "", text)
    return text.strip()


_PAGES_SECTION_RE = re.compile(
    r"Pages?\s+\d+[^\n—\-]*[—\-]\s*(?P<section>Eco-?region(?:\s+Data)?|Food|Energy|Water|"
    r"Shelter|Transportation|Conclusion|Cover(?:\s+Page)?|[A-Za-z][A-Za-z\-]{2,24})",
    re.IGNORECASE,
)


def _match_known_section(label: str, shorts: dict[str, str]) -> str | None:
    raw = (label or "").strip()
    if not raw:
        return None
    low = raw.lower()
    # Prefer longest known short that is a prefix or exact token at start.
    for key, short in sorted(shorts.items(), key=lambda kv: len(kv[0]), reverse=True):
        if low == key or low.startswith(key + " ") or low.startswith(key + "-"):
            return short
        if re.match(rf"^{re.escape(key)}\b", low):
            return short
    return None


def _guess_section_for_offset(text: str, offset: int, sections: list[str]) -> str | None:
    """Map an Insert-photos placeholder to a section title.

    Prefer ``Pages N–M — Shelter`` labels near the placeholder. Avoid matching
    bare topic lists like ``Food Energy Water Shelter Transportation``.
    """
    if not sections:
        return None
    shorts = {_section_short(s).lower(): _section_short(s) for s in sections if _section_short(s)}

    # 1) Label immediately after the placeholder (sample template style).
    after = text[offset : offset + 280]
    m = _PAGES_SECTION_RE.search(after)
    if m:
        hit = _match_known_section(m.group("section"), shorts)
        if hit:
            return hit

    # 2) Nearest preceding Pages — Section label.
    before = text[max(0, offset - 500) : offset]
    last = None
    for match in _PAGES_SECTION_RE.finditer(before):
        last = match.group("section").strip()
    if last:
        hit = _match_known_section(last, shorts)
        if hit:
            return hit

    # 3) Standalone heading line before offset (## Shelter or line == Shelter).
    window = text[:offset]
    best: tuple[int, str] | None = None
    for key, short in shorts.items():
        for match in re.finditer(
            rf"(?im)^(?:##\s+)?{re.escape(short)}\s*$",
            window,
        ):
            pos = match.start()
            if best is None or pos >= best[0]:
                best = (pos, short)
    return best[1] if best else None


def extract_figure_queries(
    text: str,
    *,
    sections: list[str] | None = None,
    community: str | None = None,
) -> list[dict[str, str]]:
    """Parse [Insert photos: ...] placeholders into Commons search queries."""
    blob = str(text or "")
    section_list = [str(s).strip() for s in (sections or []) if str(s).strip()]
    # Guess against a wider outline so Food placeholders do not stick to Shelter.
    guess_sections = list(section_list)
    for extra in (
        "Food",
        "Energy",
        "Water",
        "Shelter",
        "Transportation",
        "Eco-region Data",
        "Eco-Region",
        "Conclusion",
        "Cover",
    ):
        if not any(_section_short(s).lower() == extra.lower() for s in guess_sections):
            guess_sections.append(extra)
    allowed_shorts = {_section_short(s).lower() for s in section_list if _section_short(s)}
    community_clean = (community or "").strip()
    # Prefer a concrete place name found in text when title is generic.
    if not community_clean or community_clean.lower() in {
        "assignment",
        "untitled",
        "dissection of a sustainable community",
    }:
        place = re.search(
            r"\b(Findhorn|Arcosanti|Auroville|Dancing Rabbit|Earthaven|Twin Oaks)\b",
            blob,
            re.I,
        )
        if place:
            community_clean = place.group(1)

    queries: list[dict[str, str]] = []
    seen: set[str] = set()

    def _add(query: str, *, caption: str, section: str | None) -> None:
        q = re.sub(r"\s+", " ", query).strip(" ,;")
        if len(q) < 3:
            return
        # When the project is scoped, only keep figures for those sections.
        if allowed_shorts:
            if not section or section.lower() not in allowed_shorts:
                return
        key = q.lower()
        if key in seen:
            return
        seen.add(key)
        item: dict[str, str] = {"query": q, "caption": caption or q}
        if section:
            item["section"] = section
        queries.append(item)

    for match in _INSERT_PHOTOS_RE.finditer(blob):
        body = match.group("body") or ""
        section = _guess_section_for_offset(blob, match.start(), guess_sections)
        parts = [p.strip() for p in re.split(r",|;|/|\band\b", body, flags=re.I) if p.strip()]
        for part in parts:
            part = re.sub(r"\s+", " ", part).strip(" .")
            if len(part) < 3:
                continue
            q = f"{community_clean} {part}".strip() if community_clean else part
            _add(q, caption=part, section=section)

    # Fallback: brief asks for photos but no placeholders — one query per body section.
    if not queries and brief_requires_photos(blob):
        body_sections = [
            s
            for s in section_list
            if not re.search(r"\b(references?|bibliography|works cited|cover)\b", s, re.I)
        ]
        for section in body_sections[:6]:
            short = _section_short(section)
            q = f"{community_clean} {short}".strip() if community_clean else short
            _add(q, caption=short, section=short)

    # Cap per section so we do not flood the document.
    per_section: dict[str, int] = {}
    capped: list[dict[str, str]] = []
    for item in queries:
        sec = item.get("section") or "_general"
        if per_section.get(sec, 0) >= _MAX_IMAGES_PER_SECTION:
            continue
        per_section[sec] = per_section.get(sec, 0) + 1
        capped.append(item)
    return capped


def _safe_stem(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", (value or "").strip()).strip("-")
    return (cleaned or "figure")[:60]


def search_commons_image_url(query: str, *, session: requests.Session | None = None) -> dict[str, str] | None:
    """Return best Commons image hit for query: {title, url, mime}."""
    http = session or requests.Session()
    params = {
        "action": "query",
        "format": "json",
        "generator": "search",
        "gsrsearch": query,
        "gsrnamespace": 6,
        "gsrlimit": 8,
        "prop": "imageinfo",
        "iiprop": "url|mime|size",
        "iiurlwidth": 1280,
    }
    response = http.get(
        _COMMONS_API,
        params=params,
        headers={"User-Agent": _COMMONS_UA},
        timeout=30,
    )
    response.raise_for_status()
    pages = (response.json().get("query") or {}).get("pages") or {}
    candidates: list[dict[str, str]] = []
    for page in pages.values():
        info_list = page.get("imageinfo") or []
        if not info_list:
            continue
        info = info_list[0]
        mime = str(info.get("mime") or "").lower()
        if not mime.startswith("image/"):
            continue
        if "svg" in mime or "tif" in mime:
            continue
        url = str(info.get("thumburl") or info.get("url") or "").strip()
        if not url:
            continue
        candidates.append(
            {
                "title": str(page.get("title") or ""),
                "url": url,
                "mime": mime,
            }
        )
    return candidates[0] if candidates else None


def download_image(
    url: str,
    dest_path: Path,
    *,
    session: requests.Session | None = None,
) -> Path:
    http = session or requests.Session()
    response = http.get(url, headers={"User-Agent": _COMMONS_UA}, timeout=60)
    response.raise_for_status()
    dest_path.parent.mkdir(parents=True, exist_ok=True)
    dest_path.write_bytes(response.content)
    return dest_path


def _ext_for_mime(mime: str, url: str) -> str:
    mime = (mime or "").lower()
    if "png" in mime:
        return ".png"
    if "webp" in mime:
        return ".webp"
    if "gif" in mime:
        return ".gif"
    path_ext = Path(urlparse(url).path).suffix.lower()
    if path_ext in IMAGE_EXTENSIONS:
        return path_ext
    return ".jpg"


def fetch_commons_figures(
    *,
    project_id: str,
    dest_dir: Path,
    queries: list[dict[str, str]],
    session: requests.Session | None = None,
) -> list[dict[str, str]]:
    """Search Commons for each query and download images into dest_dir."""
    http = session or requests.Session()
    dest_dir.mkdir(parents=True, exist_ok=True)
    out: list[dict[str, str]] = []
    used_urls: set[str] = set()

    for index, item in enumerate(queries):
        query = str(item.get("query") or "").strip()
        if not query:
            continue
        try:
            hit = search_commons_image_url(query, session=http)
        except Exception:  # noqa: BLE001
            continue
        if not hit or hit["url"] in used_urls:
            # Retry with a shorter query (drop leading community once).
            parts = query.split()
            if len(parts) > 2:
                try:
                    hit = search_commons_image_url(" ".join(parts[1:]), session=http)
                except Exception:  # noqa: BLE001
                    hit = None
        if not hit or hit["url"] in used_urls:
            continue
        used_urls.add(hit["url"])
        ext = _ext_for_mime(hit.get("mime") or "", hit["url"])
        stem = _safe_stem(item.get("section") or item.get("caption") or f"figure-{index + 1}")
        filename = f"auto_{stem}_{uuid.uuid4().hex[:8]}{ext}"
        path = dest_dir / filename
        try:
            download_image(hit["url"], path, session=http)
        except Exception:  # noqa: BLE001
            continue
        if not path.is_file() or path.stat().st_size < 100:
            path.unlink(missing_ok=True)
            continue
        record: dict[str, str] = {
            "file_id": "",
            "filename": filename,
            "storage_path": str(path),
            "source": "commons",
            "source_url": hit["url"],
            "caption": str(item.get("caption") or query),
            "query": query,
        }
        if item.get("section"):
            record["section"] = str(item["section"])
        out.append(record)
    return out


def ensure_figure_markers(content: str, images: list[dict[str, Any]] | None) -> str:
    """Insert [[FIGURE:...]] into section bodies that still lack markers."""
    if not images:
        return content or ""
    text = content or ""
    # Group unused images by section (case-insensitive short title).
    by_section: dict[str, list[dict[str, Any]]] = {}
    general: list[dict[str, Any]] = []
    for item in images:
        if not isinstance(item, dict):
            continue
        fname = str(item.get("filename") or "").strip()
        if not fname or not Path(str(item.get("storage_path") or "")).is_file():
            continue
        if re.search(rf"\[\[\s*FIGURE\s*:\s*{re.escape(fname)}", text, re.I):
            continue
        sec = _section_short(str(item.get("section") or "")).lower()
        if sec:
            by_section.setdefault(sec, []).append(item)
        else:
            general.append(item)

    if not by_section and not general:
        return text

    matches = list(_HEADING_LINE_RE.finditer(text))
    if not matches:
        # No headings — append remaining figures at end.
        chunks = [text.strip()] if text.strip() else []
        for item in list(by_section.values()) + [general]:
            for img in item:
                cap = str(img.get("caption") or img.get("filename") or "").strip()
                chunks.append(f"[[FIGURE:{img['filename']}|{cap}]]")
        return "\n\n".join(chunks).strip()

    parts: list[str] = []
    if matches[0].start() > 0:
        parts.append(text[: matches[0].start()].rstrip())

    for index, match in enumerate(matches):
        start = match.start()
        end = matches[index + 1].start() if index + 1 < len(matches) else len(text)
        block = text[start:end].rstrip()
        title = match.group(1).strip()
        short = _section_short(title).lower()
        # Skip references.
        if re.search(r"\b(references?|bibliography|works cited)\b", short, re.I):
            parts.append(block)
            continue
        pending = list(by_section.pop(short, []))
        # Fuzzy section key match.
        if not pending:
            for key in list(by_section.keys()):
                if key in short or short in key:
                    pending = by_section.pop(key)
                    break
        if not pending and general and index == 0:
            pending = general
            general = []
        if not pending:
            parts.append(block)
            continue
        if FIGURE_MARKER_RE.search(block):
            # Section already has some markers — only add missing filenames.
            missing = [
                img
                for img in pending
                if not re.search(rf"\[\[\s*FIGURE\s*:\s*{re.escape(str(img['filename']))}", block, re.I)
            ]
            pending = missing
        if not pending:
            parts.append(block)
            continue
        # Insert after heading line + first body paragraph when possible.
        lines = block.split("\n", 1)
        heading = lines[0]
        rest = lines[1] if len(lines) > 1 else ""
        markers = "\n\n".join(
            f"[[FIGURE:{img['filename']}|{str(img.get('caption') or img['filename']).strip()}]]"
            for img in pending
        )
        if rest.strip():
            # After first paragraph block.
            paras = re.split(r"\n\s*\n", rest.lstrip(), maxsplit=1)
            if len(paras) == 2:
                new_block = f"{heading}\n\n{paras[0].strip()}\n\n{markers}\n\n{paras[1].strip()}"
            else:
                new_block = f"{heading}\n\n{paras[0].strip()}\n\n{markers}"
        else:
            new_block = f"{heading}\n\n{markers}"
        parts.append(new_block.strip())

    # Leftover images (section title mismatch) → append before references if any.
    leftovers: list[dict[str, Any]] = []
    for items in by_section.values():
        leftovers.extend(items)
    leftovers.extend(general)
    body = "\n\n".join(p for p in parts if p).strip()
    if leftovers:
        markers = "\n\n".join(
            f"[[FIGURE:{img['filename']}|{str(img.get('caption') or img['filename']).strip()}]]"
            for img in leftovers
        )
        ref_match = re.search(r"(?im)^##\s+(References|Works Cited|Bibliography)\s*$", body)
        if ref_match:
            body = body[: ref_match.start()].rstrip() + "\n\n" + markers + "\n\n" + body[ref_match.start() :]
        else:
            body = body + "\n\n" + markers
    return body.strip()


def merge_image_lists(*lists: list[dict[str, Any]] | None) -> list[dict[str, str]]:
    """Dedupe image dicts by resolved storage path."""
    out: list[dict[str, str]] = []
    seen: set[str] = set()
    for group in lists:
        for item in group or []:
            if not isinstance(item, dict):
                continue
            path = Path(str(item.get("storage_path") or ""))
            if not path.is_file():
                continue
            key = str(path.resolve())
            if key in seen:
                continue
            seen.add(key)
            record = {str(k): str(v) for k, v in item.items() if v is not None}
            record["filename"] = Path(str(item.get("filename") or path.name)).name
            record["storage_path"] = str(path)
            out.append(record)
    return out
