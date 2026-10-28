"""Title page for assignment DOCX output (title, brief identifiers, required declaration)."""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt


def title_page_wanted(requirement_json: dict[str, Any]) -> bool:
    page = requirement_json.get("title_page")
    if isinstance(page, dict) and (page.get("required") or page.get("declaration")):
        return True
    for section in requirement_json.get("required_sections") or []:
        name = str(section).lower()
        if "title page" in name or "cover page" in name:
            return True
    return False


def add_title_page(docx_bytes: bytes, *, title: str, requirement_json: dict[str, Any]) -> bytes:
    """Prepend a title page; drops the body's leading title paragraph so it is not printed twice."""
    page = requirement_json.get("title_page") if isinstance(requirement_json.get("title_page"), dict) else {}
    lines = [
        str(v).strip()
        for v in (page.get("lines") or [])
        if str(v).strip() and str(v).strip().lower() != title.strip().lower()
    ]
    declaration = str(page.get("declaration") or "").strip()

    doc = Document(io.BytesIO(docx_bytes))
    body = doc.element.body
    first_para = doc.paragraphs[0] if doc.paragraphs else None
    if first_para is not None and first_para.text.strip() == title.strip():
        body.remove(first_para._p)

    new_paras = []

    def add(text: str, *, bold: bool = False, size: float | None = None, space_before: float = 0,
            align=WD_ALIGN_PARAGRAPH.CENTER):
        para = doc.add_paragraph()
        para.alignment = align
        para.paragraph_format.first_line_indent = Pt(0)
        para.paragraph_format.space_before = Pt(space_before)
        para.paragraph_format.space_after = Pt(6)
        if text:
            run = para.add_run(text)
            run.bold = bold
            if size:
                run.font.size = Pt(size)
        new_paras.append(para)
        return para

    add(title, bold=True, size=18, space_before=150)
    for index, line in enumerate(lines):
        add(line, space_before=24 if index == 0 else 0)
    if declaration:
        add(declaration, space_before=150, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    breaker = add("")
    breaker.add_run().add_break(WD_BREAK.PAGE)

    anchor = body[0]
    for para in new_paras:
        element = para._p
        body.remove(element)
        anchor.addprevious(element)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
