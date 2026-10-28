"""Markdown pipe tables -> real Word tables.

Formatter V2 rebuilds the document paragraph by paragraph and has no table
support, so tables travel through it as ``[[TABLE:n]]`` placeholder paragraphs
and are swapped in afterwards.
"""

from __future__ import annotations

import io
import re

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

_SEPARATOR_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)*\|?$")
_APPENDIX_HEADING_RE = re.compile(r"^Appendix [A-Z]\s*:")


def _cells(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def extract_markdown_tables(content: str) -> tuple[str, list[list[list[str]]]]:
    """Replace each pipe-table block with ``[[TABLE:n]]``; return the parsed rows."""
    tables: list[list[list[str]]] = []
    out: list[str] = []
    for block in re.split(r"(\n\s*\n)", content or ""):
        lines = [ln.strip() for ln in block.strip().split("\n") if ln.strip()]
        if (
            len(lines) >= 3
            and all(ln.startswith("|") for ln in lines)
            and _SEPARATOR_RE.match(lines[1])
        ):
            width = len(_cells(lines[0]))
            rows = [_cells(lines[0])] + [_cells(ln) for ln in lines[2:]]
            tables.append([(r + [""] * width)[:width] for r in rows])
            out.append(f"[[TABLE:{len(tables) - 1}]]")
        else:
            out.append(block)
    return "".join(out), tables


def insert_word_tables(docx_bytes: bytes, tables: list[list[list[str]]]) -> bytes:
    if not tables:
        return docx_bytes
    doc = Document(io.BytesIO(docx_bytes))
    _unify_appendix_headings(doc)
    for paragraph in list(doc.paragraphs):
        match = re.fullmatch(r"\[\[TABLE:(\d+)\]\]", paragraph.text.strip())
        if not match or int(match.group(1)) >= len(tables):
            continue
        rows = tables[int(match.group(1))]
        table = doc.add_table(rows=len(rows), cols=len(rows[0]))
        _set_borders(table)
        for r, row in enumerate(rows):
            for c, value in enumerate(row):
                cell_par = table.cell(r, c).paragraphs[0]
                fmt = cell_par.paragraph_format
                fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
                fmt.space_before = fmt.space_after = 0
                fmt.first_line_indent = 0
                run = cell_par.add_run(value)
                run.bold = r == 0
        if len(rows) > 1:
            _repeat_header(table.rows[0])
        paragraph._p.addprevious(table._tbl)
        paragraph._p.getparent().remove(paragraph._p)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _unify_appendix_headings(doc) -> None:
    """Formatter V2 only recognises the first "Appendix A: ..." after the References."""
    headings = [p for p in doc.paragraphs if _APPENDIX_HEADING_RE.match(p.text.strip())]
    if not headings:
        return
    style = next((p.style for p in headings if p.style.name != "Normal"), headings[0].style)
    for paragraph in headings:
        paragraph.style = style
        paragraph.paragraph_format.page_break_before = True


def _set_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tbl_pr.append(borders)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)
