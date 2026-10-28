"""Turn the Editor's document HTML (contenteditable output) into a .docx file."""

from __future__ import annotations

import io
import re
from typing import Any, Iterable

from bs4 import BeautifulSoup, Comment, NavigableString, Tag
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

MAX_HTML_CHARS = 2_000_000
MAX_TITLE_CHARS = 200

_ALIGN = {
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}
_HEADINGS = {"h1": 1, "h2": 2, "h3": 3, "h4": 3, "h5": 3, "h6": 3}
_BLOCKS = {"p", "div", "ul", "ol", "li", "blockquote", "pre", *_HEADINGS}
_SKIP = {"script", "style", "button", "template"}
_INDENT_STEP = Inches(0.5)


class EditorExportError(ValueError):
    pass


def export_filename(title: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", str(title or ""), flags=re.UNICODE).strip()
    slug = re.sub(r"[\s_-]+", "-", slug)[:60].strip("-")
    return f"{slug or 'document'}.docx"


def html_to_docx(title: str, html: str) -> bytes:
    html = str(html or "")
    if len(html) > MAX_HTML_CHARS:
        raise EditorExportError("The document is too large to export.")
    doc = Document()
    _base_styles(doc)
    doc.core_properties.title = str(title or "")[:MAX_TITLE_CHARS]
    _render_children(doc, BeautifulSoup(html, "html.parser"), indent=0)
    if not doc.paragraphs:
        doc.add_paragraph()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _base_styles(doc: Any) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for level, size in ((1, 18), (2, 15), (3, 13)):
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Times New Roman"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)


def _is_block(node: Any) -> bool:
    return isinstance(node, Tag) and node.name in _BLOCKS


def _render_children(doc: Any, parent: Any, *, indent: int) -> None:
    """Blocks become paragraphs; loose inline content between blocks becomes its own paragraph."""
    loose: list[Any] = []

    def flush() -> None:
        if any(_has_text(n) for n in loose):
            _paragraph(doc, loose, indent=indent)
        loose.clear()

    for child in parent.children:
        if _is_block(child):
            flush()
            _render_block(doc, child, indent=indent)
        else:
            loose.append(child)
    flush()


def _render_block(doc: Any, el: Tag, *, indent: int) -> None:
    name = el.name
    if name in ("ul", "ol"):
        _render_list(doc, el, ordered=name == "ol", level=1, indent=indent)
    elif name == "blockquote":
        _render_children(doc, el, indent=indent + 1)
    elif name in ("div", "li") and any(_is_block(c) for c in el.children):
        _render_children(doc, el, indent=indent)
    elif name in _HEADINGS:
        _paragraph(doc, el.children, indent=indent, heading=_HEADINGS[name], source=el)
    else:
        _paragraph(doc, el.children, indent=indent, source=el)


def _render_list(doc: Any, el: Tag, *, ordered: bool, level: int, indent: int) -> None:
    style = ("List Number" if ordered else "List Bullet") + (f" {min(level, 3)}" if level > 1 else "")
    for li in el.find_all("li", recursive=False):
        inline = [c for c in li.children if not (isinstance(c, Tag) and c.name in ("ul", "ol"))]
        _paragraph(doc, inline, indent=indent, style=style, source=li)
        for sub in li.find_all(["ul", "ol"], recursive=False):
            _render_list(doc, sub, ordered=sub.name == "ol", level=level + 1, indent=indent)


def _paragraph(
    doc: Any,
    nodes: Iterable[Any],
    *,
    indent: int,
    heading: int | None = None,
    style: str | None = None,
    source: Tag | None = None,
) -> None:
    p = doc.add_heading(level=heading) if heading else doc.add_paragraph(style=style)
    if indent:
        p.paragraph_format.left_indent = _INDENT_STEP * indent
    align = _align(source)
    if align is not None:
        p.alignment = align
    _inline(p, nodes, {})
    if p.runs:
        p.runs[0].text = p.runs[0].text.lstrip()
        p.runs[-1].text = p.runs[-1].text.rstrip()


def _align(el: Tag | None) -> Any:
    if el is None:
        return None
    m = re.search(r"text-align\s*:\s*(left|center|right|justify)", el.get("style") or "", re.I)
    value = (m.group(1) if m else el.get("align") or "").lower()
    return _ALIGN.get(value)


def _inline(p: Any, nodes: Iterable[Any], fmt: dict[str, bool]) -> None:
    for node in nodes:
        if isinstance(node, Comment):
            continue
        if isinstance(node, NavigableString):
            text = re.sub(r"\s+", " ", str(node))
            if text:
                _run(p, text, fmt)
            continue
        if not isinstance(node, Tag) or node.name in _SKIP:
            continue
        if node.name == "br":
            p.add_run().add_break()
            continue
        _inline(p, node.children, _format(node, fmt))


def _format(el: Tag, fmt: dict[str, bool]) -> dict[str, bool]:
    out = dict(fmt)
    name = el.name
    css = (el.get("style") or "").lower()
    if name in ("b", "strong") or re.search(r"font-weight\s*:\s*(bold|[6-9]00)", css):
        out["bold"] = True
    if name in ("i", "em") or "font-style: italic" in css or "font-style:italic" in css:
        out["italic"] = True
    if name == "u" or "underline" in css:
        out["underline"] = True
    if name in ("s", "strike", "del") or "line-through" in css:
        out["strike"] = True
    if name == "sup":
        out["superscript"] = True
    if name == "sub":
        out["subscript"] = True
    return out


def _run(p: Any, text: str, fmt: dict[str, bool]) -> None:
    run = p.add_run(text)
    run.bold = fmt.get("bold") or None
    run.italic = fmt.get("italic") or None
    run.underline = fmt.get("underline") or None
    if fmt.get("strike"):
        run.font.strike = True
    if fmt.get("superscript"):
        run.font.superscript = True
    elif fmt.get("subscript"):
        run.font.subscript = True


def _has_text(node: Any) -> bool:
    if isinstance(node, Comment):
        return False
    if isinstance(node, NavigableString):
        return bool(str(node).strip())
    return isinstance(node, Tag) and node.name not in _SKIP and bool(node.get_text().strip())
