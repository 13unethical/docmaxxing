"""AI presentation generator: deck normalization, PPTX rendering, and API wiring."""

from __future__ import annotations

import io
from unittest.mock import MagicMock

import pytest
from pptx import Presentation

import app as app_module
from services import presentation_generator as pg
from services import presentation_images as pimg
from services.presentation_themes import DEFAULT_THEME, THEMES


@pytest.fixture(autouse=True)
def no_photo_network(monkeypatch, tmp_path):
    monkeypatch.setenv("PRESENTATION_IMAGE_DIR", str(tmp_path / "images"))
    monkeypatch.setattr(pg, "attach_images", lambda deck, layouts: 0)


def _cached_image(image_id: str = "a" * 20, size=(1200, 800)) -> dict[str, str]:
    from PIL import Image

    folder = pimg.image_dir()
    folder.mkdir(parents=True, exist_ok=True)
    Image.new("RGB", size, (40, 90, 160)).save(folder / f"{image_id}.jpg", format="JPEG")
    return {"id": image_id, "credit": "Photo: Jane Doe / Wikimedia Commons", "source": "https://commons.example/x"}

SAMPLE_DECK = {
    "title": "Renewable Energy in Central Asia",
    "subtitle": "Potential, barriers, and policy",
    "slides": [
        {"layout": "bullets", "title": "Why it matters", "bullets": ["Rising demand", "Aging grids"], "notes": "Intro."},
        {
            "layout": "two_column",
            "title": "Solar vs wind",
            "left": {"heading": "Solar", "bullets": ["High irradiance"]},
            "right": {"heading": "Wind", "bullets": ["Steppe corridors"]},
        },
        {"layout": "section", "title": "Policy", "subtitle": "What governments can do"},
        {"layout": "closing", "title": "Key takeaways", "bullets": ["Huge potential", "Needs investment"]},
    ],
}


def test_normalize_deck_bounds_and_fallbacks():
    raw = {
        "title": "  T  ",
        "slides": [
            {"layout": "weird", "title": "A", "bullets": ["x"] * 20},
            {"layout": "two_column", "title": "B", "left": {}, "right": {}, "bullets": ["y"]},
            {"layout": "bullets", "title": "", "bullets": []},
            "not a slide",
        ],
    }
    deck = pg.normalize_deck(raw)
    assert deck["title"] == "T"
    assert [s["layout"] for s in deck["slides"]] == ["bullets", "bullets"]
    assert len(deck["slides"][0]["bullets"]) == pg.MAX_BULLETS


def test_normalize_deck_rejects_empty():
    with pytest.raises(pg.PresentationError):
        pg.normalize_deck({"slides": []})
    with pytest.raises(pg.PresentationError):
        pg.normalize_deck("nope")


def test_build_pptx_renders_all_layouts_with_notes():
    data = pg.build_pptx(SAMPLE_DECK)
    prs = Presentation(io.BytesIO(data))
    slides = list(prs.slides)
    assert len(slides) == 5
    text = "\n".join(
        shape.text_frame.text for slide in slides for shape in slide.shapes if shape.has_text_frame
    )
    for expected in ("Renewable Energy in Central Asia", "Solar", "Steppe corridors", "Key takeaways"):
        assert expected in text
    assert slides[1].notes_slide.notes_text_frame.text == "Intro."


def test_generate_deck_prompt_and_clamp(monkeypatch):
    captured = {}

    def fake_generate_json(**kwargs):
        captured.update(kwargs)
        return SAMPLE_DECK, {"model": "test"}

    monkeypatch.setattr(pg, "generate_json", fake_generate_json)
    deck, _diag = pg.generate_deck(topic="Energy", slide_count=99, language="ru")
    assert len(deck["slides"]) == 4
    assert f"exactly {pg.MAX_SLIDES - 1} slides" in captured["user_prompt"]
    assert "Russian" in captured["user_prompt"]


def test_normalize_deck_theme_falls_back_to_default():
    assert pg.normalize_deck({**SAMPLE_DECK, "theme": "midnight"})["theme"] == "midnight"
    assert pg.normalize_deck({**SAMPLE_DECK, "theme": "nope"})["theme"] == DEFAULT_THEME
    assert pg.normalize_deck(SAMPLE_DECK)["theme"] == DEFAULT_THEME


@pytest.mark.parametrize("theme_id", list(THEMES))
def test_build_pptx_every_theme(theme_id):
    prs = Presentation(io.BytesIO(pg.build_pptx({**SAMPLE_DECK, "theme": theme_id})))
    assert len(prs.slides) == 5
    theme = THEMES[theme_id]
    cover_bg = prs.slides[0].shapes[0]
    assert str(cover_bg.fill.fore_color.rgb) == theme.cover_bg.upper()


def test_generate_deck_applies_theme(monkeypatch):
    monkeypatch.setattr(pg, "generate_json", lambda **_: (dict(SAMPLE_DECK), {"model": "test"}))
    deck, _diag = pg.generate_deck(topic="Energy", theme="pitch")
    assert deck["theme"] == "pitch"


def test_generate_deck_raises_on_model_failure(monkeypatch):
    monkeypatch.setattr(pg, "generate_json", lambda **_: (None, {"error_message": "boom"}))
    with pytest.raises(pg.PresentationError):
        pg.generate_deck(topic="Energy")


def test_deck_filename():
    assert pg.deck_filename({"title": "AI: The Future?"}) == "AI-The-Future.pptx"
    assert pg.deck_filename({"title": "!!!"}) == "presentation.pptx"


@pytest.fixture()
def client(monkeypatch):
    monkeypatch.setattr("services.economy.auth.email_verification_gate", lambda **_: None)
    monkeypatch.setattr("services.claude_client.claude_enabled", lambda: True)
    app_module.app.config["TESTING"] = True
    return app_module.app.test_client()


def test_presentation_page_renders(client):
    res = client.get("/presentation")
    assert res.status_code == 200
    html = res.get_data(as_text=True)
    assert "pres_generate_btn" in html
    assert "pres_theme_grid" in html
    assert '"id": "midnight"' in html or '"id":"midnight"' in html
    assert "SOON" not in html


def test_api_generate_charges_and_returns_deck(client, monkeypatch):
    charge = MagicMock(return_value=(7, {"balance": 70}))
    refund = MagicMock()
    monkeypatch.setattr(app_module, "_charge_current_user", charge)
    monkeypatch.setattr(app_module, "_refund_safe", refund)
    monkeypatch.setattr(pg, "generate_deck", lambda **_: (pg.normalize_deck(SAMPLE_DECK), {"model": "m"}))

    res = client.post("/api/presentation/generate", json={"topic": "Energy", "slide_count": 5})
    body = res.get_json()
    assert res.status_code == 200
    assert body["success"] is True
    assert body["balance"] == 70
    assert len(body["deck"]["slides"]) == 4
    charge.assert_called_once()
    assert charge.call_args.args[0] == "presentation"
    refund.assert_not_called()


def test_api_generate_refunds_on_failure(client, monkeypatch):
    monkeypatch.setattr(app_module, "_charge_current_user", MagicMock(return_value=(7, {"balance": 70})))
    refund = MagicMock()
    monkeypatch.setattr(app_module, "_refund_safe", refund)

    def boom(**_):
        raise pg.PresentationError("model down")

    monkeypatch.setattr(pg, "generate_deck", boom)
    res = client.post("/api/presentation/generate", json={"topic": "Energy"})
    assert res.status_code == 502
    refund.assert_called_once()
    assert refund.call_args.args[2] == "presentation"


def test_api_generate_requires_topic(client, monkeypatch):
    charge = MagicMock()
    monkeypatch.setattr(app_module, "_charge_current_user", charge)
    res = client.post("/api/presentation/generate", json={"topic": "  "})
    assert res.status_code == 400
    charge.assert_not_called()


def test_api_export_returns_pptx(client):
    res = client.post("/api/presentation/export", json={"deck": SAMPLE_DECK})
    assert res.status_code == 200
    assert "presentationml" in res.headers["Content-Type"]
    assert "Renewable-Energy-in-Central-Asia.pptx" in res.headers["Content-Disposition"]
    assert len(Presentation(io.BytesIO(res.data)).slides) == 5


def test_api_export_rejects_bad_deck(client):
    res = client.post("/api/presentation/export", json={"deck": {"slides": []}})
    assert res.status_code == 400


VISUAL_SLIDES = [
    {"layout": "image", "title": "In the classroom", "bullets": ["Laptops", "Feedback"], "image_query": "students laptops"},
    {
        "layout": "cards",
        "title": "Benefits",
        "items": [{"heading": f"Benefit {i}", "text": f"Explained {i}."} for i in range(1, 5)],
    },
    {
        "layout": "timeline",
        "title": "Rollout",
        "items": [{"heading": "Audit", "text": "Find tasks."}, {"heading": "Pilot", "text": "One group."}, {"heading": "Scale", "text": "All."}],
    },
    {"layout": "section", "title": "Teachers", "subtitle": "New role", "image_query": "teacher helping student"},
    {"layout": "closing", "title": "Takeaways", "bullets": ["Done"]},
]


def test_normalize_visual_layouts_and_fallbacks():
    deck = pg.normalize_deck(
        {
            "title": "T",
            "image_query": "  campus   lawn ",
            "slides": VISUAL_SLIDES
            + [{"layout": "cards", "title": "Too few", "items": [{"heading": "Only", "text": "one"}]}],
        }
    )
    layouts = [s["layout"] for s in deck["slides"]]
    assert layouts == ["image", "cards", "timeline", "section", "closing", "bullets"]
    assert deck["image_query"] == "campus lawn"
    assert len(deck["slides"][1]["items"]) == 4
    assert deck["slides"][-1]["bullets"] == ["Only: one"]


def test_normalize_rejects_untrusted_image_ids():
    deck = pg.normalize_deck(
        {
            **SAMPLE_DECK,
            "cover_image": {"id": "../../etc/passwd", "credit": "x"},
            "slides": [{"layout": "image", "title": "A", "bullets": ["b"], "image": {"id": "b" * 20, "credit": "c"}}],
        }
    )
    assert deck["cover_image"] is None
    assert deck["slides"][0]["image"] == {"id": "b" * 20, "url": f"/api/presentation/image/{'b' * 20}", "credit": "c", "source": ""}


def test_build_pptx_places_photos_and_credits():
    image = _cached_image()
    slides = [dict(s) for s in VISUAL_SLIDES]
    slides[0]["image"] = image
    slides[3]["image"] = image
    prs = Presentation(io.BytesIO(pg.build_pptx({**SAMPLE_DECK, "cover_image": image, "slides": slides})))
    pictures = [sum(1 for sh in slide.shapes if sh.shape_type == 13) for slide in prs.slides]
    assert pictures == [1, 1, 0, 0, 1, 0]
    assert "Jane Doe" in prs.slides[0].notes_slide.notes_text_frame.text
    assert "Jane Doe" in prs.slides[1].notes_slide.notes_text_frame.text
    text = "\n".join(sh.text_frame.text for sh in prs.slides[2].shapes if sh.has_text_frame)
    assert "Benefit 4" in text and "04" in text


def test_build_pptx_lists_notes_master_for_keynote():
    import zipfile

    data = pg.build_pptx(SAMPLE_DECK)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("ppt/presentation.xml").decode()
    assert "<p:notesMasterIdLst><p:notesMasterId " in xml
    assert xml.index("<p:sldMasterIdLst>") < xml.index("<p:notesMasterIdLst>") < xml.index("<p:sldIdLst>")
    assert 'type="screen4x3"' not in xml


def test_image_slide_without_cached_photo_renders_as_bullets():
    deck = {**SAMPLE_DECK, "slides": [{**VISUAL_SLIDES[0], "image": {"id": "c" * 20}}]}
    prs = Presentation(io.BytesIO(pg.build_pptx(deck)))
    assert not any(sh.shape_type == 13 for sh in prs.slides[1].shapes)
    assert "Laptops" in "\n".join(sh.text_frame.text for sh in prs.slides[1].shapes if sh.has_text_frame)


def test_generate_deck_attaches_photos(monkeypatch):
    raw = {"title": "T", "image_query": "campus", "slides": VISUAL_SLIDES}
    monkeypatch.setattr(pg, "generate_json", lambda **_: (dict(raw), {"model": "test"}))
    calls = []

    def fake_attach(deck, layouts):
        calls.append(layouts)
        deck["slides"][3]["image"] = {"id": "d" * 20}
        return 1

    monkeypatch.setattr(pg, "attach_images", fake_attach)
    deck, diag = pg.generate_deck(topic="AI")
    assert calls == [pg.PHOTO_LAYOUTS]
    assert diag["images_found"] == 1
    # No photo was found for the image slide, so it degrades to a plain list.
    assert deck["slides"][0]["layout"] == "bullets"


def test_generate_deck_without_photos(monkeypatch):
    captured = {}

    def fake_generate_json(**kwargs):
        captured.update(kwargs)
        return {"title": "T", "slides": VISUAL_SLIDES}, {"model": "test"}

    monkeypatch.setattr(pg, "generate_json", fake_generate_json)
    monkeypatch.setattr(pg, "attach_images", MagicMock(side_effect=AssertionError("should not fetch")))
    _deck, diag = pg.generate_deck(topic="AI", images=False)
    assert "Photos are turned off" in captured["user_prompt"]
    assert diag["images_found"] == 0


class _FakeResponse:
    def __init__(self, *, payload=None, content=b""):
        self._payload = payload or {}
        self.content = content

    def raise_for_status(self):
        return None

    def json(self):
        return self._payload


def test_fetch_image_downloads_downscales_and_caches(monkeypatch):
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (3000, 2000), (200, 10, 10)).save(buf, format="PNG")
    monkeypatch.setattr(
        pimg, "search_images", lambda query, **_: [{"url": "https://img.example/a.png", "credit": "Photo: A", "source": "s"}]
    )
    http = MagicMock()
    http.get.return_value = _FakeResponse(content=buf.getvalue())

    first = pimg.fetch_image("students on laptops", session=http)
    second = pimg.fetch_image("students on laptops", session=http)
    assert first == second
    assert first["url"] == f"/api/presentation/image/{first['id']}"
    path = pimg.image_path(first["id"])
    with Image.open(path) as im:
        assert im.format == "JPEG" and max(im.size) == 1600
    assert http.get.call_count == 1


def test_commons_search_prefers_landscape_photos():
    def page(index, title, mime="image/jpeg", w=2000, h=1000):
        return {
            "index": index,
            "title": title,
            "imageinfo": [{"mime": mime, "width": w, "height": h, "thumburl": f"https://u/{index}.jpg", "extmetadata": {}}],
        }

    pages = {
        "1": page(1, "File:Students vertical shot.jpg", w=800, h=1600),
        "2": page(2, "File:School logo.jpg"),
        "3": page(3, "File:Diagram.png", mime="image/png"),
        "4": page(4, "File:Students in class.jpg"),
        "5": page(5, "File:Tiny.jpg", w=300, h=200),
    }
    http = MagicMock()
    http.get.return_value = _FakeResponse(payload={"query": {"pages": pages}})
    hits = pimg._search_commons("students", http)
    assert [h["url"] for h in hits] == ["https://u/4.jpg", "https://u/1.jpg"]


def test_commons_search_requires_title_to_match_the_subject():
    def page(index, title):
        info = {"mime": "image/jpeg", "width": 2000, "height": 1200, "thumburl": f"https://u/{index}.jpg", "extmetadata": {}}
        return {"index": index, "title": title, "imageinfo": [info]}

    pages = {
        "1": page(1, "File:NASA students fly on DC-8.jpg"),
        "2": page(2, "File:Minister of Education addresses students.jpg"),
        "3": page(3, "File:Student in the library, 1981.jpg"),
    }
    http = MagicMock()
    http.get.return_value = _FakeResponse(payload={"query": {"pages": pages}})
    hits = pimg._search_commons("student reading in library", http)
    assert [h["url"] for h in hits] == ["https://u/3.jpg"]


def test_api_image_serves_cached_file_and_rejects_bad_ids(client):
    image = _cached_image("e" * 20)
    res = client.get(f"/api/presentation/image/{image['id']}")
    assert res.status_code == 200
    assert res.headers["Content-Type"] == "image/jpeg"
    assert client.get("/api/presentation/image/..%2F..%2Fapp.py").status_code == 404
    assert client.get("/api/presentation/image/" + "f" * 20).status_code == 404


def test_api_generate_passes_photo_flag(client, monkeypatch):
    monkeypatch.setattr(app_module, "_charge_current_user", MagicMock(return_value=(7, {"balance": 70})))
    seen = {}

    def fake_generate_deck(**kwargs):
        seen.update(kwargs)
        return pg.normalize_deck(SAMPLE_DECK), {"model": "m"}

    monkeypatch.setattr(pg, "generate_deck", fake_generate_deck)
    client.post("/api/presentation/generate", json={"topic": "Energy", "images": False})
    assert seen["images"] is False
    client.post("/api/presentation/generate", json={"topic": "Energy"})
    assert seen["images"] is True


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Create a professional 8-slide presentation", 8),
        ("Сделай презентацию на 10 слайдов", 10),
        ("a 3 slide deck", None),
        ("Top 100 slides of all time", None),
        ("No count here", None),
    ],
)
def test_requested_slide_count(text, expected):
    assert pg.requested_slide_count(text) == expected


def test_generate_deck_prefers_slide_count_written_in_request(monkeypatch):
    captured = {}

    def fake_generate_json(**kwargs):
        captured.update(kwargs)
        return SAMPLE_DECK, {"model": "test"}

    monkeypatch.setattr(pg, "generate_json", fake_generate_json)
    pg.generate_deck(topic="Create a professional 10-slide presentation about startups", slide_count=8)
    assert "exactly 9 slides" in captured["user_prompt"]


def test_presenter_on_cover():
    deck = pg.normalize_deck({**SAMPLE_DECK, "presenter": "Daniel Karimov · BrightPath Student Lab"})
    assert deck["presenter"] == "Daniel Karimov · BrightPath Student Lab"
    prs = Presentation(io.BytesIO(pg.build_pptx(deck)))
    cover_text = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert "Daniel Karimov · BrightPath Student Lab" in cover_text


@pytest.mark.parametrize(
    "text, expected",
    [
        ("Presentation duration: 7–10 minutes.", 8.5),
        ("a 5 minute talk", 5),
        ("доклад на 10 минут", 10),
        ("no duration", None),
        ("90 minutes lecture", None),
    ],
)
def test_requested_minutes(text, expected):
    assert pg.requested_minutes(text) == expected


def test_generate_deck_sizes_notes_to_duration(monkeypatch):
    captured = {}

    def fake_generate_json(**kwargs):
        captured.update(kwargs)
        return SAMPLE_DECK, {"model": "test"}

    monkeypatch.setattr(pg, "generate_json", fake_generate_json)
    pg.generate_deck(topic="8-slide deck. Duration: 7–10 minutes.")
    assert "about 130 words of speaker notes" in captured["user_prompt"]


def test_shortened_retry_is_judged_against_the_full_subject():
    pages = {
        "1": {
            "index": 1,
            "title": "File:Al-Nawfara Coffee Shop, Damascus.jpg",
            "imageinfo": [{"mime": "image/jpeg", "width": 2000, "height": 1200, "thumburl": "https://u/1.jpg", "extmetadata": {}}],
        }
    }
    http = MagicMock()
    http.get.return_value = _FakeResponse(payload={"query": {"pages": pages}})
    assert pimg._search_commons("coffee shop", http)
    assert pimg._search_commons("coffee shop", http, "young entrepreneur campus coffee shop") == []


def test_themes_cover_every_design_and_ship_sample_photos():
    from pathlib import Path

    from services.presentation_themes import CATEGORIES, DESIGNS

    samples = Path(__file__).resolve().parent.parent / "static" / "presentation" / "samples"
    categories = {key for key, _label in CATEGORIES}
    assert {t.design for t in THEMES.values()} == set(DESIGNS)
    for theme in THEMES.values():
        assert theme.category in categories
        assert theme.title_accent in ("", "lead", "tail")
        assert (samples / f"{theme.sample_image}.jpg").is_file(), theme.id


def test_accent_split():
    from services.presentation_themes import accent_split

    assert accent_split("Opening a Neighborhood Coffee Shop", "lead") == [
        ("Opening ", True),
        ("a Neighborhood Coffee Shop", False),
    ]
    assert accent_split("Mountain Tourism Development Plan", "tail") == [
        ("Mountain Tourism Development ", False),
        ("Plan", True),
    ]
    assert accent_split("Strategy", "tail") == [("Strategy", False)]
    assert accent_split("Two words here", "") == [("Two words here", False)]


@pytest.mark.parametrize("design", ["classic", "magazine", "diagonal", "editorial", "fullbleed", "lineart"])
def test_every_design_renders_photos_with_preset_shapes_only(design):
    """Quick Look / iOS drop custom geometry, so designs must stick to preset shapes."""
    import zipfile

    theme_id = next(t.id for t in THEMES.values() if t.design == design)
    image = _cached_image()
    slides = [dict(s) for s in VISUAL_SLIDES]
    slides[0]["image"] = image
    slides[3]["image"] = image
    data = pg.build_pptx({**SAMPLE_DECK, "theme": theme_id, "cover_image": image, "slides": slides})
    archive = zipfile.ZipFile(io.BytesIO(data))
    slide_xml = [archive.read(n).decode() for n in archive.namelist() if n.startswith("ppt/slides/slide")]
    assert not any("<a:custGeom>" in xml for xml in slide_xml)
    prs = Presentation(io.BytesIO(data))
    assert len(prs.slides) == 6
    assert any(sh.shape_type == 13 for sh in prs.slides[0].shapes)
    cover_text = "\n".join(sh.text_frame.text for sh in prs.slides[0].shapes if sh.has_text_frame)
    assert THEMES[theme_id].heading_upper or "Renewable Energy" in cover_text


def test_accent_colors_part_of_the_cover_title():
    prs = Presentation(io.BytesIO(pg.build_pptx({**SAMPLE_DECK, "theme": "sky"})))
    theme = THEMES["sky"]
    runs = [
        run
        for sh in prs.slides[0].shapes
        if sh.has_text_frame
        for para in sh.text_frame.paragraphs
        for run in para.runs
        if run.text.strip() in ("RENEWABLE ENERGY IN CENTRAL", "ASIA")
    ]
    colors = {run.text.strip(): str(run.font.color.rgb) for run in runs}
    assert colors == {"RENEWABLE ENERGY IN CENTRAL": theme.cover_ink, "ASIA": theme.cover_accent}


def test_lineart_cover_crops_photo_to_a_circle():
    from PIL import Image

    image = _cached_image()
    prs = Presentation(io.BytesIO(pg.build_pptx({**SAMPLE_DECK, "theme": "paper", "cover_image": image})))
    pic = next(sh for sh in prs.slides[0].shapes if sh.shape_type == 13)
    assert pic.width == pic.height
    with Image.open(io.BytesIO(pic.image.blob)) as im:
        assert im.size[0] == im.size[1]
        assert im.getpixel((2, 2)) == tuple(int(THEMES["paper"].cover_bg[i : i + 2], 16) for i in (0, 2, 4))


# --- attached files + web research -------------------------------------------------


def _capture_generate_json(monkeypatch, response=None):
    captured = {}

    def fake_generate_json(**kwargs):
        captured.update(kwargs)
        return dict(response or SAMPLE_DECK), {"model": "test"}

    monkeypatch.setattr(pg, "generate_json", fake_generate_json)
    return captured


def test_clean_attachments_bounds_count_and_total_text():
    raw = [{"name": f"f{i}.pdf", "text": "x" * 25000} for i in range(5)] + ["junk"]
    files = pg.clean_attachments(raw)
    assert len(files) == pg.MAX_ATTACHMENTS
    assert sum(len(f["text"]) for f in files) == pg.MAX_ATTACHMENT_CHARS
    assert pg.clean_attachments([{"name": 'a"b', "text": "  "}]) == []
    assert pg.clean_attachments([{"name": 'a"b', "text": "hi"}])[0]["name"] == "a'b"
    assert pg.clean_attachments("nope") == []


def test_generate_deck_builds_from_attached_brief_without_topic(monkeypatch):
    captured = _capture_generate_json(monkeypatch)
    brief = "Assignment: a 10-slide presentation on water scarcity in Uzbekistan, 7-10 minutes."
    deck, diag = pg.generate_deck(topic="", attachments=[{"name": "brief.pdf", "text": brief}])
    prompt = captured["user_prompt"]
    assert "take it from the attached files" in prompt
    assert '<file name="brief.pdf">' in prompt and brief in prompt
    assert "exactly 9 slides" in prompt
    assert "words of speaker notes" in prompt
    assert "research" not in diag
    assert deck["slides"]


def test_generate_deck_requires_topic_or_file():
    with pytest.raises(pg.PresentationError):
        pg.generate_deck(topic="  ", attachments=[{"name": "empty.txt", "text": "   "}])


def test_generate_deck_with_research_adds_notes_and_sources_slide(monkeypatch):
    captured = _capture_generate_json(monkeypatch, {**SAMPLE_DECK, "sources_title": "Источники"})
    seen = {}

    def fake_research(topic, context=""):
        seen.update(topic=topic, context=context)
        sources = [
            {"url": "https://www.iea.org/reports/renewables", "title": "Renewables 2026"},
            {"url": "https://irena.org/data", "title": ""},
        ]
        return "Solar capacity in Uzbekistan reached 4 GW in 2026.", sources, {"web_searches": 3}

    monkeypatch.setattr(pg, "research_topic", fake_research)
    deck, diag = pg.generate_deck(topic="Renewables", research=True, details="Focus on Uzbekistan")
    assert seen == {"topic": "Renewables", "context": "Focus on Uzbekistan"}
    assert "<research>\nSolar capacity in Uzbekistan reached 4 GW in 2026.\n</research>" in captured["user_prompt"]
    assert diag["research"] == "ok" and diag["web_searches"] == 3
    last = deck["slides"][-1]
    assert last["layout"] == "bullets"
    assert last["title"] == "Источники"
    assert last["bullets"] == ["Renewables 2026 — iea.org", "irena.org"]
    assert last["notes"] == "1) https://www.iea.org/reports/renewables · 2) https://irena.org/data"
    assert len(deck["slides"]) == len(SAMPLE_DECK["slides"]) + 1


def test_generate_deck_research_failure_still_builds_deck(monkeypatch):
    captured = _capture_generate_json(monkeypatch)
    monkeypatch.setattr(pg, "research_topic", lambda topic, context="": (None, [], {"web_searches": 0}))
    deck, diag = pg.generate_deck(topic="Renewables", research=True)
    assert diag["research"] == "failed"
    assert "<research>" not in captured["user_prompt"]
    assert len(deck["slides"]) == len(SAMPLE_DECK["slides"])


def test_api_generate_accepts_files_without_topic_and_skips_research(client, monkeypatch):
    charge = MagicMock(return_value=(7, {"balance": 50}))
    refund = MagicMock()
    monkeypatch.setattr(app_module, "_charge_current_user", charge)
    monkeypatch.setattr(app_module, "_refund_safe", refund)
    seen = {}

    def fake_generate_deck(**kwargs):
        seen.update(kwargs)
        return pg.normalize_deck(SAMPLE_DECK), {"model": "m"}

    monkeypatch.setattr(pg, "generate_deck", fake_generate_deck)
    res = client.post(
        "/api/presentation/generate",
        json={"topic": "", "files": [{"name": "brief.docx", "text": "Brief text"}], "research": True},
    )
    body = res.get_json()
    assert res.status_code == 200, body
    expected = app_module.feature_cost("presentation")
    assert charge.call_args.args[:2] == ("presentation", expected)
    assert body["credits_used"] == expected
    assert seen["attachments"] == [{"name": "brief.docx", "text": "Brief text"}]
    assert seen["research"] is False
    assert seen["slide_count"] is None
    refund.assert_not_called()


def test_api_generate_researches_bare_topic_automatically(client, monkeypatch):
    charge = MagicMock(return_value=(7, {"balance": 50}))
    monkeypatch.setattr(app_module, "_charge_current_user", charge)
    seen = {}

    def fake_generate_deck(**kwargs):
        seen.update(kwargs)
        return pg.normalize_deck(SAMPLE_DECK), {"model": "m", "research": "ok"}

    monkeypatch.setattr(pg, "generate_deck", fake_generate_deck)
    res = client.post("/api/presentation/generate", json={"topic": "Сделай презентацию про ИИ в медицине"})
    body = res.get_json()
    expected = app_module.feature_cost("presentation") + app_module.feature_cost("presentation_research")
    assert charge.call_args.args[:2] == ("presentation", expected)
    assert body["credits_used"] == expected
    assert body["research"] == "ok"
    assert seen["research"] is True
    assert seen["language"] == "auto"


def test_api_generate_refunds_research_surcharge_when_research_fails(client, monkeypatch):
    monkeypatch.setattr(app_module, "_charge_current_user", MagicMock(return_value=(7, {"balance": 50})))
    refund = MagicMock()
    monkeypatch.setattr(app_module, "_refund_safe", refund)
    monkeypatch.setattr(app_module.wallet, "get_balance", lambda _uid: 70)
    monkeypatch.setattr(
        pg, "generate_deck", lambda **_: (pg.normalize_deck(SAMPLE_DECK), {"model": "m", "research": "failed"})
    )
    res = client.post("/api/presentation/generate", json={"topic": "Energy"})
    body = res.get_json()
    research_cost = app_module.feature_cost("presentation_research")
    refund.assert_called_once_with(7, research_cost, "presentation_research")
    assert body["credits_used"] == app_module.feature_cost("presentation")
    assert body["balance"] == 70
    assert body["research"] == "failed"


def test_api_generate_with_pasted_material_charges_base_price(client, monkeypatch):
    charge = MagicMock(return_value=(7, {"balance": 70}))
    monkeypatch.setattr(app_module, "_charge_current_user", charge)
    monkeypatch.setattr(pg, "generate_deck", lambda **_: (pg.normalize_deck(SAMPLE_DECK), {"model": "m"}))
    client.post("/api/presentation/generate", json={"topic": "Energy", "details": "My notes", "research": True})
    assert charge.call_args.args[1] == app_module.feature_cost("presentation")


def test_api_generate_treats_long_pasted_text_as_material(client, monkeypatch):
    charge = MagicMock(return_value=(7, {"balance": 70}))
    monkeypatch.setattr(app_module, "_charge_current_user", charge)
    seen = {}

    def fake_generate_deck(**kwargs):
        seen.update(kwargs)
        return pg.normalize_deck(SAMPLE_DECK), {"model": "m"}

    monkeypatch.setattr(pg, "generate_deck", fake_generate_deck)
    notes = "Lecture notes.\n" + "Photosynthesis converts light into chemical energy. " * 40
    client.post("/api/presentation/generate", json={"topic": notes})
    assert charge.call_args.args[1] == app_module.feature_cost("presentation")
    assert seen["research"] is False


def test_generate_deck_keeps_pasted_material_layout(monkeypatch):
    captured = _capture_generate_json(monkeypatch)
    pg.generate_deck(topic="Make slides from my notes:\n- point one\n- point two")
    assert "Make slides from my notes:\n- point one\n- point two" in captured["user_prompt"]


def test_generate_deck_picks_slide_count_automatically(monkeypatch):
    captured = _capture_generate_json(monkeypatch)
    pg.generate_deck(topic="История Узбекистана")
    prompt = captured["user_prompt"]
    assert "exactly" not in prompt
    assert "Choose how many slides the content needs: 7-11 slides" in prompt
    assert "the language the user wrote the topic or instructions in" in prompt

    pg.generate_deck(topic="Доклад на 10 минут про климат")
    assert "exactly 10 slides" in captured["user_prompt"]
    assert "words of speaker notes" in captured["user_prompt"]


def test_api_extract_reads_text_and_pptx_files(client):
    res = client.post(
        "/api/presentation/extract",
        data={"file": (io.BytesIO("Тема: водные ресурсы".encode()), "brief.txt")},
        content_type="multipart/form-data",
    )
    body = res.get_json()
    assert res.status_code == 200, body
    assert body["text"] == "Тема: водные ресурсы" and body["name"] == "brief.txt"
    assert body["truncated"] is False

    pptx = pg.build_pptx(SAMPLE_DECK)
    res = client.post(
        "/api/presentation/extract",
        data={"file": (io.BytesIO(pptx), "old-deck.pptx")},
        content_type="multipart/form-data",
    )
    body = res.get_json()
    assert res.status_code == 200, body
    assert "Slide 2:" in body["text"] and "Why it matters" in body["text"]
    assert "Notes: Intro." in body["text"]


def test_api_extract_rejects_bad_files(client):
    res = client.post(
        "/api/presentation/extract",
        data={"file": (io.BytesIO(b"MZ"), "virus.exe")},
        content_type="multipart/form-data",
    )
    assert res.status_code == 400
    res = client.post(
        "/api/presentation/extract",
        data={"file": (io.BytesIO(b"not a zip"), "broken.pptx")},
        content_type="multipart/form-data",
    )
    assert res.status_code == 400
    assert client.post("/api/presentation/extract").status_code == 400


def test_presentation_page_shows_attach_without_manual_pickers(client):
    html = client.get("/presentation").get_data(as_text=True)
    assert 'id="pres_attach_btn"' in html
    for removed in ("pres_research_btn", "pres_slide_count", "pres_language", "pres-mode-chip", "pres_details"):
        assert removed not in html
    assert ".pptx" in html
    assert f'data-research-cost="{app_module.feature_cost("presentation_research")}"' in html
