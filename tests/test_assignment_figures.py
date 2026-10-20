"""FIGURE markers and image embedding for assignment delivery."""

from __future__ import annotations

import io
from pathlib import Path
from unittest.mock import MagicMock

from docx import Document
from PIL import Image

from services.assignment_figures import (
    brief_requires_photos,
    collect_assignment_images,
    embed_figures_in_docx_bytes,
    ensure_figure_markers,
    extract_figure_queries,
    fetch_commons_figures,
    protect_figure_markers,
    restore_figure_markers,
)
from services.assignment_project.models import ProjectFile, ProjectFileType
from services.delivery_engine.packager import _build_docx_bytes


def _tiny_png(path: Path) -> None:
    img = Image.new("RGB", (32, 24), color=(40, 120, 200))
    img.save(path, format="PNG")


def test_collect_assignment_images_skips_non_images(tmp_path):
    png = tmp_path / "shelter.png"
    _tiny_png(png)
    pdf = tmp_path / "brief.pdf"
    pdf.write_bytes(b"%PDF-1.4")
    files = [
        ProjectFile(
            id="1",
            project_id="p",
            file_type=ProjectFileType.ADDITIONAL_FILE,
            filename="shelter.png",
            original_filename="shelter.png",
            storage_path=str(png),
        ),
        ProjectFile(
            id="2",
            project_id="p",
            file_type=ProjectFileType.ASSIGNMENT_BRIEF,
            filename="brief.pdf",
            original_filename="brief.pdf",
            storage_path=str(pdf),
        ),
    ]
    images = collect_assignment_images(files)
    assert len(images) == 1
    assert images[0]["filename"] == "shelter.png"


def test_protect_and_restore_figure_markers():
    text = "Body prose.\n\n[[FIGURE:shelter.png|Findhorn cabin]]\n\nMore prose."
    protected, figures = protect_figure_markers(text)
    assert "FIGURE:shelter.png" not in protected
    assert "[[[FIGURE_MARK_0]]]" in protected
    restored = restore_figure_markers(protected, figures)
    assert "[[FIGURE:shelter.png|Findhorn cabin]]" in restored


def test_embed_figures_in_docx_bytes(tmp_path):
    png = tmp_path / "transport.png"
    _tiny_png(png)
    doc = Document()
    doc.add_paragraph("Before.")
    doc.add_paragraph("[[FIGURE:transport.png|Shared bikes]]")
    doc.add_paragraph("After.")
    buf = io.BytesIO()
    doc.save(buf)
    out = embed_figures_in_docx_bytes(
        buf.getvalue(),
        [{"filename": "transport.png", "storage_path": str(png)}],
    )
    result = Document(io.BytesIO(out))
    joined = "\n".join(p.text for p in result.paragraphs)
    assert "[[FIGURE:" not in joined
    assert any(p._element.xpath(".//a:blip") for p in result.paragraphs)


def test_embed_figures_matches_underscore_stripped_names(tmp_path):
    """Markdown italic cleanup can turn auto_Shelter_abc.jpg into autoShelterabc.jpg."""
    png = tmp_path / "auto_Shelter_abc123.jpg"
    _tiny_png(png)
    doc = Document()
    doc.add_paragraph("[[FIGURE:autoShelterabc123.jpg|cabin]]")
    buf = io.BytesIO()
    doc.save(buf)
    out = embed_figures_in_docx_bytes(
        buf.getvalue(),
        [{"filename": "auto_Shelter_abc123.jpg", "storage_path": str(png)}],
    )
    result = Document(io.BytesIO(out))
    assert "[[FIGURE:" not in "\n".join(p.text for p in result.paragraphs)
    assert any(p._element.xpath(".//a:blip") for p in result.paragraphs)


def test_build_docx_bytes_embeds_figures(tmp_path):
    png = tmp_path / "a.png"
    _tiny_png(png)
    raw = _build_docx_bytes(
        "Essay",
        "Intro text.\n\n[[FIGURE:a.png|Caption]]\n\nClosing.",
        images=[{"filename": "a.png", "storage_path": str(png)}],
    )
    doc = Document(io.BytesIO(raw))
    assert any(p._element.xpath(".//a:blip") for p in doc.paragraphs)


def test_extract_figure_queries_from_sample_placeholders():
    text = (
        "Include photos that illustrate the data.\n"
        "[ Insert photos: whisky-barrel houses, eco-houses, earthships ] "
        "Pages 9–10 of 14 — Shelter\n"
        "Transportation text here.\n"
        "[ Insert photos: EV charging point, bike paths ] "
        "Pages 11–12 of 14 — Transportation\n"
        "[ Insert photos: Cullerne Garden, EarthShare produce ] "
        "Pages 3–4 of 14 — Food\n"
    )
    assert brief_requires_photos(text) is True
    queries = extract_figure_queries(
        text,
        sections=["Shelter", "Transportation", "References"],
        community="Findhorn",
    )
    assert queries
    captions = " ".join(q["caption"].lower() for q in queries)
    assert "whisky" in captions or "eco-houses" in captions
    assert "cullerne" not in captions  # Food scoped out
    assert any(q.get("section") == "Shelter" for q in queries)
    assert any(q.get("section") == "Transportation" for q in queries)
    shelter = [q for q in queries if q.get("section") == "Shelter"]
    assert len(shelter) <= 2


def test_ensure_figure_markers_injects_into_section(tmp_path):
    png = tmp_path / "auto_shelter.png"
    _tiny_png(png)
    content = (
        "## Shelter\n\nBarrel houses are common.\n\nMore analysis here.\n\n"
        "## Transportation\n\nCarshare works.\n"
    )
    images = [
        {
            "filename": "auto_shelter.png",
            "storage_path": str(png),
            "caption": "whisky-barrel houses",
            "section": "Shelter",
        }
    ]
    out = ensure_figure_markers(content, images)
    assert "[[FIGURE:auto_shelter.png|whisky-barrel houses]]" in out
    assert out.index("## Shelter") < out.index("[[FIGURE:auto_shelter.png")
    assert out.index("[[FIGURE:auto_shelter.png") < out.index("## Transportation")
    assert ensure_figure_markers(out, images) == out


def test_fetch_commons_figures_downloads(tmp_path):
    png_bytes = io.BytesIO()
    Image.new("RGB", (40, 30), color=(10, 80, 40)).save(png_bytes, format="PNG")
    payload = png_bytes.getvalue()

    class FakeResp:
        def __init__(self, data=None, content=b"", status=200):
            self._data = data or {}
            self.content = content
            self.status_code = status

        def raise_for_status(self):
            if self.status_code >= 400:
                raise RuntimeError("http error")

        def json(self):
            return self._data

    session = MagicMock()

    def fake_get(url, params=None, headers=None, timeout=None):
        del headers, timeout
        if "api.php" in str(url) or (params and params.get("action") == "query"):
            return FakeResp(
                {
                    "query": {
                        "pages": {
                            "1": {
                                "title": "File:Test.jpg",
                                "imageinfo": [
                                    {
                                        "url": "https://example.com/test.jpg",
                                        "thumburl": "https://example.com/test.jpg",
                                        "mime": "image/jpeg",
                                    }
                                ],
                            }
                        }
                    }
                }
            )
        return FakeResp(content=payload)

    session.get.side_effect = fake_get
    out = fetch_commons_figures(
        project_id="p1",
        dest_dir=tmp_path / "auto",
        queries=[
            {
                "query": "Findhorn whisky barrel",
                "caption": "barrel",
                "section": "Shelter",
            }
        ],
        session=session,
    )
    assert len(out) == 1
    assert Path(out[0]["storage_path"]).is_file()
    assert out[0]["source"] == "commons"
    assert out[0]["section"] == "Shelter"
