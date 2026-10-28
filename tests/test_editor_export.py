import io

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from services.editor_export import EditorExportError, export_filename, html_to_docx


def _doc(html: str, title: str = "Essay"):
    return Document(io.BytesIO(html_to_docx(title, html)))


def test_headings_paragraphs_and_inline_formatting():
    doc = _doc(
        "<h1>Title</h1><h2>Intro</h2>"
        "<p style=\"text-align: center;\">Plain <b>bold</b> <i>italic</i> <u>under</u> H<sub>2</sub>O x<sup>2</sup></p>"
    )
    paras = doc.paragraphs
    assert [p.style.name for p in paras] == ["Heading 1", "Heading 2", "Normal"]
    assert paras[2].text == "Plain bold italic under H2O x2"
    assert paras[2].alignment == WD_ALIGN_PARAGRAPH.CENTER
    runs = {r.text: r for r in paras[2].runs}
    assert runs["bold"].bold and runs["italic"].italic and runs["under"].underline
    assert runs["2"].font.subscript or runs["2"].font.superscript
    assert doc.core_properties.title == "Essay"


def test_lists_blockquote_and_loose_text():
    doc = _doc("Loose text<ul><li>One<ul><li>Nested</li></ul></li><li>Two</li></ul>"
               "<ol><li>First</li></ol><blockquote><p>Quoted</p></blockquote><p>   </p>")
    got = [(p.style.name, p.text) for p in doc.paragraphs]
    assert got[:5] == [
        ("Normal", "Loose text"),
        ("List Bullet", "One"),
        ("List Bullet 2", "Nested"),
        ("List Bullet", "Two"),
        ("List Number", "First"),
    ]
    quoted = doc.paragraphs[5]
    assert quoted.text == "Quoted" and quoted.paragraph_format.left_indent > 0


def test_empty_document_still_exports():
    assert len(_doc("").paragraphs) == 1


def test_too_large_is_rejected():
    with pytest.raises(EditorExportError):
        html_to_docx("Big", "x" * 2_000_001)


def test_filename():
    assert export_filename("My essay: draft #2") == "My-essay-draft-2.docx"
    assert export_filename("") == "document.docx"
