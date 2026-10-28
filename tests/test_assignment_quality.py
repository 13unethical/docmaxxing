"""Assignment quality fixes: headings, brief materials, citations, title page, appendices, grade gate."""

from __future__ import annotations

import io
import json
from types import SimpleNamespace

import docx

from services import assignment_citations as cites
from services.assignment_appendices import (
    _parse_appendices,
    appendices_markdown,
    appendices_wanted,
    generate_appendices,
    link_appendices,
)
from services.assignment_formatting import AssignmentFormatEngine
from services.assignment_formatting.tables import extract_markdown_tables, insert_word_tables
from services.assignment_formatting.title_page import add_title_page, title_page_wanted
from services.assignment_spec import build_assignment_spec, run_grade_gate
from services.assignment_spec.titles import budget_for, section_heading, section_name
from services.assignment_spec.validate import validate_draft_against_spec
from services.research_engine.parsed_documents import writer_source_material
from services.writer_engine.llm_writer import materials_system_blocks

KAPLAN_SECTIONS = [
    "Section 1: Executive Summary (Approx. 200 words)",
    "Section 2: Introduction to EcoNova Solutions (Approx. 200 words)",
    "Section 7: Conclusion (Approx. 150 words)",
    "Section 8: References",
    "Optional Appendices",
]
DECLARATION = (
    "I confirm that this assignment is my own work. Where I have referred to academic sources, "
    "I have provided in-text citations and included the sources in the final reference list."
)


def _kaplan_requirement(**extra) -> dict:
    return {
        "title": "Simulated Report",
        "word_count": 2000,
        "citation_style": "Harvard",
        "required_sections": list(KAPLAN_SECTIONS),
        "formatting": {"font_family": "Arial", "font_size": 11, "line_spacing": 1.5},
        **extra,
    }


def _body(words: int) -> str:
    sentence = "EcoNova expands its solar installation business in the south of England. "
    return (sentence * (words // 11 + 1)).strip()


# --- headings and budgets -------------------------------------------------


def test_section_heading_keeps_descriptive_name():
    assert section_heading("Section 1: Executive Summary (Approx. 200 words)") == "Section 1: Executive Summary"
    assert section_heading("Section 2: Introduction to EcoNova Solutions") == "Section 2: Introduction to EcoNova Solutions"
    assert section_heading("Introduction: set the scene for the report and explain the purpose of the analysis") == "Introduction"
    assert section_name("Section 3: External Environment Analysis") == "External Environment Analysis"


def test_budget_lookup_does_not_confuse_section_1_and_10():
    budgets = {"Section 1": 200, "Section 10": 90}
    assert budget_for("Section 1: Executive Summary", budgets) == 200
    assert budget_for("Section 10: Reflection", budgets) == 90


def test_spec_keeps_full_titles_and_skips_appendices():
    spec = build_assignment_spec(_kaplan_requirement(), project_id="p")
    by_title = {s.title: s for s in spec.sections}
    assert by_title["Section 1: Executive Summary"].target_words == 200
    assert by_title["Section 8: References"].writable is False
    assert by_title["Optional Appendices"].mandatory is False


def test_validator_matches_numbered_references_and_ignores_optional_appendices():
    spec = build_assignment_spec(_kaplan_requirement(), project_id="p")
    content = (
        f"## Section 1: Executive Summary\n\n{_body(200)}\n\n"
        f"## Section 2: Introduction to EcoNova Solutions\n\n{_body(200)}\n\n"
        f"## Section 7: Conclusion\n\n{_body(150)}\n\n"
        "## References\n\nBarney, J. (1991) 'Firm resources', Journal of Management, 17(1), pp. 99-120."
    )
    result = validate_draft_against_spec(content=content, spec=spec)
    assert result.missing_sections == []


# --- brief materials for the writer ----------------------------------------


def test_writer_material_puts_brief_first_and_keeps_line_breaks(tmp_path):
    brief = tmp_path / "brief.txt"
    brief.write_text("EcoNova Solutions\nFounded 2016 in Winchester\n\nRevenue £4.2m", encoding="utf-8")
    notes = tmp_path / "notes.txt"
    notes.write_text("Lecture notes on PESTLE", encoding="utf-8")
    files = [
        SimpleNamespace(file_type="lecture_notes", storage_path=str(notes), original_filename="notes.txt"),
        SimpleNamespace(file_type="assignment_brief", storage_path=str(brief), original_filename="brief.txt"),
    ]
    text = writer_source_material(files, note="Use UK spelling")
    assert text.startswith("[student_note]\nUse UK spelling")
    assert text.index("brief.txt") < text.index("notes.txt")
    assert "Founded 2016 in Winchester\n\nRevenue £4.2m" in text


def test_materials_travel_in_a_cached_system_block():
    blocks = materials_system_blocks("CASE: EcoNova")
    assert blocks[0]["cache_control"] == {"type": "ephemeral"}
    assert "<materials>\nCASE: EcoNova\n</materials>" in blocks[0]["text"]


# --- citations --------------------------------------------------------------


def test_parenthetical_citation_groups_are_extracted():
    text = "Demand is rising (DESNZ, 2024; Porter, 1980). Barney (1991) argues otherwise (see Table 1, 2020)."
    found = cites._extract_claimed_cites(text)
    assert "Barney 1991" in found
    assert "DESNZ 2024" in found
    assert "Porter 1980" in found
    assert not any(q.lower().startswith("see") for q in found)


def test_crossref_hit_must_match_author_and_year():
    good = {"authors": ["Porter, M.E."], "year": "1980", "label": "Porter (1980)"}
    wrong_author = {"authors": ["Rotunno, L."], "year": "1980", "label": "Rotunno (1980)"}
    anon = {"authors": [], "year": "1980", "label": "Anon. (1980)"}
    assert cites._work_matches_query("Porter 1980", good)
    assert not cites._work_matches_query("Porter 1980", wrong_author)
    assert not cites._work_matches_query("Porter 1980", anon)
    assert not cites._work_matches_query("Porter 1995", good)


def test_organisation_acronym_matches_full_name():
    work = {"authors": ["Department for Energy Security and Net Zero,"], "year": "2024", "label": "x"}
    assert cites._work_matches_query("DESNZ 2024", work)


def test_web_verification_requires_a_searched_source(monkeypatch):
    payload = {
        "items": [
            {"cite": "Porter 1980", "found": True, "title": "Competitive Strategy", "authors": ["Porter, M.E."],
             "year": "1980", "publisher": "Free Press", "url": "https://www.simonandschuster.com/x"},
            {"cite": "Turner 2021", "found": True, "title": "Made up", "authors": ["Turner, A."],
             "year": "2021", "url": "https://invented.example.com/paper"},
        ]
    }
    sources = [{"url": "https://simonandschuster.com/books/competitive-strategy"}]
    monkeypatch.setattr(
        "services.claude_client.generate_with_web_search",
        lambda **_kw: (json.dumps(payload), sources, {}),
    )
    confirmed = cites._web_verify_citations(["Porter 1980", "Turner 2021"], style="Harvard")
    assert [q for q, _ in confirmed] == ["Porter 1980"]
    reference = confirmed[0][1]["reference"]
    assert "Free Press" in reference and "Accessed:" in reference


# --- title page ---------------------------------------------------------------


def _docx_with(*paragraphs: str) -> bytes:
    doc = docx.Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def test_title_page_has_declaration_and_replaces_duplicate_title():
    req = {"title_page": {"required": True, "declaration": DECLARATION, "lines": ["IYO Business", "Simulated Report"]}}
    assert title_page_wanted(req)
    out = add_title_page(_docx_with("Simulated Report", "Body text."), title="Simulated Report", requirement_json=req)
    texts = [p.text for p in docx.Document(io.BytesIO(out)).paragraphs]
    assert texts[:3] == ["Simulated Report", "IYO Business", DECLARATION]
    assert texts.count("Simulated Report") == 1
    assert texts[-1] == "Body text."


def test_title_page_not_added_without_request():
    assert not title_page_wanted({"title_page": {}, "required_sections": ["Introduction"]})
    assert title_page_wanted({"required_sections": ["Title page", "Introduction"]})


# --- appendices and tables --------------------------------------------------

APPENDICES = [
    {"title": "PESTLE Analysis of EcoNova Solutions", "keyword": "PESTLE", "intro": "External factors.",
     "headers": ["Factor", "Evidence", "Impact"],
     "rows": [["Political", "UK net zero target", "Opportunity"], ["Economic", "High interest rates", "Threat"]]},
    {"title": "SWOT Analysis", "keyword": "SWOT", "intro": "",
     "headers": ["Strengths", "Weaknesses"], "rows": [["Brand", "Cash flow"]]},
]


def test_appendices_wanted_only_when_brief_allows_them():
    assert appendices_wanted({"required_sections": ["Optional Appendices"]})
    assert appendices_wanted({}, "Section 8: References\nOptional Appendices")
    assert not appendices_wanted({"required_sections": ["Introduction"]}, "Write 2000 words.")


def test_generate_appendices_parses_and_bounds_the_model_output():
    raw = "```json\n" + json.dumps({"appendices": APPENDICES + [{"title": "No table", "headers": ["x"], "rows": []}]}) + "\n```"
    calls = []
    result = generate_appendices(materials="CASE", draft_content="REPORT", call=lambda p, m: calls.append((p, m)) or raw)
    assert [a["title"] for a in result] == ["PESTLE Analysis of EcoNova Solutions", "SWOT Analysis"]
    assert "REPORT" in calls[0][0] and calls[0][1] == "CASE"
    assert _parse_appendices("not json") == []


def test_link_appendices_points_first_body_mention():
    content = "## Section 3: External\n\nA PESTLE review shows support. More PESTLE text.\n\n## SWOT\n\nThe SWOT shows gaps."
    linked = link_appendices(content, APPENDICES)
    assert "A PESTLE review shows support (see Appendix A)." in linked
    assert linked.count("Appendix A") == 1
    assert "The SWOT shows gaps (see Appendix B)." in linked
    assert "## SWOT\n" in linked


def test_markdown_tables_become_word_tables():
    md = appendices_markdown(APPENDICES)
    text, tables = extract_markdown_tables(f"Intro.\n\n{md}")
    assert "[[TABLE:0]]" in text and "[[TABLE:1]]" in text
    assert tables[0][0] == ["Factor", "Evidence", "Impact"]
    out = insert_word_tables(_docx_with("Appendix A: PESTLE", "[[TABLE:0]]", "[[TABLE:1]]"), tables)
    doc = docx.Document(io.BytesIO(out))
    assert len(doc.tables) == 2
    assert doc.tables[0].rows[2].cells[0].text == "Economic"
    assert all("[[TABLE" not in p.text for p in doc.paragraphs)


def test_format_engine_renders_title_page_and_appendix_tables(tmp_path, monkeypatch):
    monkeypatch.setenv("PROJECT_STORAGE_DIR", str(tmp_path))
    req = _kaplan_requirement(title_page={"required": True, "declaration": DECLARATION, "lines": ["IYO Business"]})
    content = link_appendices(
        "## Section 1: Executive Summary\n\nThe PESTLE review shows strong policy support.\n\n"
        "## References\n\nBarney, J. (1991) 'Firm resources', Journal of Management, 17(1), pp. 99-120.",
        APPENDICES,
    )
    out = AssignmentFormatEngine().format_draft(
        draft={"title": "Simulated Report", "content": content},
        requirement_json=req,
        project_id="p1",
        appendices=appendices_markdown(APPENDICES),
    )
    doc = docx.Document(out["path"])
    texts = [p.text for p in doc.paragraphs]
    assert texts[0] == "Simulated Report" and DECLARATION in texts
    headings = [p for p in doc.paragraphs if p.text.startswith("Appendix ")]
    assert [p.text.split(":")[0] for p in headings] == ["Appendix A", "Appendix B"]
    assert len({p.style.name for p in headings}) == 1
    assert all(p.paragraph_format.page_break_before for p in headings)
    assert len(doc.tables) == 2
    assert "Appendix A:" not in out["plain_text"]


# --- grade gate ---------------------------------------------------------------


def test_grade_gate_never_appends_template_sentences():
    spec = build_assignment_spec(_kaplan_requirement(), project_id="p")
    content = (
        "## Section 1: Executive Summary\n\nToo short.\n\n"
        f"## Section 2: Introduction to EcoNova Solutions\n\n{_body(200)}\n\n"
        f"## Section 7: Conclusion\n\n{_body(150)}\n\n"
        "## References\n\nBarney, J. (1991) 'Firm resources', Journal of Management, 17(1), pp. 99-120."
    )
    result = run_grade_gate(content=content, spec=spec, llm_repair=None)
    for phrase in ("lecture and seminar", "entrepreneurial judgement", "historical concept", "major choice"):
        assert phrase not in result.content


def test_grade_gate_trims_whole_sentences():
    spec = build_assignment_spec(_kaplan_requirement(), project_id="p")
    content = (
        f"## Section 1: Executive Summary\n\n{_body(400)}\n\n"
        f"## Section 2: Introduction to EcoNova Solutions\n\n{_body(200)}\n\n"
        f"## Section 7: Conclusion\n\n{_body(150)}\n\n"
        "## References\n\nBarney, J. (1991) 'Firm resources', Journal of Management, 17(1), pp. 99-120."
    )
    result = run_grade_gate(content=content, spec=spec, llm_repair=None)
    summary = result.content.split("## Section 2")[0]
    body = summary.split("\n\n", 1)[1].strip()
    assert body.endswith("England.")
    assert len(body.split()) <= 220


def test_trim_drops_signposting_but_keeps_ending_and_citations():
    from services.assignment_spec.grade_gate import _trim_to_sentences

    body = (
        "EcoNova installs solar panels for homes in the south of England and has grown steadily. "
        "This section will now discuss the external environment in detail. "
        "Energy prices rose sharply, which raised demand for solar systems (DESNZ, 2024). "
        "Demand for solar systems rose sharply as energy prices rose across the south of England. "
        "EcoNova should therefore prioritise residential installations."
    )
    out = _trim_to_sentences(body, 40)
    assert "will now discuss" not in out
    assert "(DESNZ, 2024)" in out
    assert out.startswith("EcoNova installs")
    assert out.endswith("prioritise residential installations.")


# --- headings and incremental humanization -----------------------------------


def test_long_heading_is_not_split_into_heading_and_body():
    from services.humanizer_engine.heading_utils import normalize_markdown_headings

    text = "## Introduction to EcoNova Solutions\n\nEcoNova Solutions is a firm."
    assert normalize_markdown_headings(text) == text
    glued = "It will compete. ## Conclusion EcoNova should expand into commercial clients across the south."
    assert "\n\n## Conclusion\n\nEcoNova should" in normalize_markdown_headings(glued)


def test_rehumanization_sends_changed_sections_whole_in_one_call():
    from services.humanizer_engine.incremental import humanize_new_sentences

    previous = (
        "## Intro\n\nOld humanized sentence that already exists in the draft. Another old sentence stays put.\n\n"
        "## Analysis\n\nThe analysis paragraph was humanized earlier and nobody changed it at all.\n\n"
        "## Conclusion\n\nThe conclusion was humanized earlier too and it should be left alone."
    )
    body = previous.replace(
        "Another old sentence stays put.",
        "A brand new sentence was added by the grade gate repair step. Another old sentence stays put.",
    ).replace(
        "it should be left alone.", "it should be left alone. A second new sentence now closes the report firmly."
    )
    sent: list[str] = []

    def fake(text: str) -> str:
        sent.append(text)
        return text.replace("sentence", "line")

    out = humanize_new_sentences(body, previous, fake)
    assert len(sent) == 1
    assert "## Intro" in sent[0] and "## Conclusion" in sent[0] and "Analysis" not in sent[0]
    assert "Old humanized line that already exists" in out
    assert "The analysis paragraph was humanized earlier and nobody changed it at all." in out
    assert out.startswith("## Intro")


def test_rehumanization_skips_untouched_document():
    from services.humanizer_engine.incremental import humanize_new_sentences

    text = "## Intro\n\nNothing changed in this section since it was humanized last time."
    assert humanize_new_sentences(text, text, lambda _t: "SHOULD NOT BE CALLED") == text


def test_meaning_repair_applies_small_edits_and_never_touches_references():
    from services.humanizer_engine.fidelity import repair_humanizer_meaning

    original = (
        "## Analysis\n\nIts resources are valuable, rare and inimitable (Barney, 1991). It has about 55 employees.\n\n"
        "## References\n\nBarney, J. (1991) 'Firm resources', Journal of Management, 17(1), pp. 99-120."
    )
    humanized = (
        "## Analysis\n\nIts resources are valuable, rare and imitable (Barney, 1991). The firm has less than 55 "
        "staff.\n\n## References\n\nBarney, J. (1991) 'Firm resources', Journal of Management, 17(1), pp. 99-120."
    )
    seen: list[str] = []

    def fake(_system, user):
        seen.append(user)
        return {"edits": [
            {"find": "rare and imitable", "replace": "rare and inimitable"},
            {"find": "less than 55", "replace": "about 55"},
            {"find": "Its resources are valuable, rare and imitable (Barney, 1991).",
             "replace": "Its resources are valuable (Smith, 2020)."},
            {"find": "Journal of Management", "replace": "Journal of Business"},
        ]}

    fixed, edits = repair_humanizer_meaning(original, humanized, call=fake)
    assert "rare and inimitable (Barney, 1991)" in fixed
    assert "about 55 staff" in fixed
    assert "Smith" not in fixed and "Journal of Management" in fixed
    assert len(edits) == 2
    assert all("References" not in u for u in seen)


# --- citation accuracy -------------------------------------------------------


def test_crossref_hit_needs_every_co_author_and_a_topic_link():
    query = "Worthington and Britton 2015"
    context = "EcoNova operates in a dynamic business environment (Worthington and Britton, 2015)."
    akkadian = {"authors": ["Worthington, M."], "title": "Akkadian", "year": "2015"}
    off_topic = {"authors": ["Worthington, I.", "Britton, C."], "title": "Akkadian grammar", "year": "2015"}
    textbook = {"authors": ["Worthington, I.", "Britton, C."], "title": "The Business Environment", "year": "2015"}
    assert not cites._work_matches_query(query, akkadian, context=context)
    assert not cites._work_matches_query(query, off_topic, context=context)
    assert cites._work_matches_query(query, textbook, context=context)


def test_single_author_cite_rejects_multi_author_work_and_off_topic_hits():
    context = "Rivalry is assessed with the five forces of Porter (1980), which shape competitive strategy."
    medical = {"authors": ["Porter, A.", "Porter, J.M.T."], "title": "Anglo-French contrasts in medical practice",
               "year": "1980"}
    assert not cites._work_matches_query("Porter 1980", medical, context=context)
    space = {"authors": ["Barney, R."], "title": "Contamination Control Program for the Cosmic Background Explorer",
             "year": "1991"}
    rbv = "Under the resource-based view, valuable and rare resources give sustained advantage (Barney, 1991)."
    assert not cites._work_matches_query("Barney 1991", space, context=rbv)
    good = {"authors": ["Barney, J."], "title": "Firm Resources and Sustained Competitive Advantage", "year": "1991"}
    assert cites._work_matches_query("Barney 1991", good, context=rbv)
    assert cites._extract_claimed_cites("Porter's (1980) five forces apply.") == ["Porter 1980"]


def test_summaries_and_reviews_of_cited_authors_are_rejected():
    names = frozenset({"elkington", "ansoff"})
    review = {
        "authors": ["Jeurissen, R."],
        "title": "John Elkington, Cannibals with Forks: The Triple Bottom Line",
        "year": "2000",
    }
    assert not cites._work_matches_query("triple bottom line business", review, other_cited_names=names)
    original = {"authors": ["Elkington, J."], "title": "Cannibals with Forks", "year": "1997"}
    assert cites._work_matches_query("Elkington 1997", original, other_cited_names=names)


def test_reconcile_may_not_swap_authors():
    before = "Growth options follow Ansoff (1957)."
    assert cites._introduces_new_cited_authors(before, "Growth options follow Mulders (2019).")
    assert not cites._introduces_new_cited_authors(before, "Growth options follow (Ansoff, 1957).")


def test_cited_name_spelling_follows_the_reference_list():
    content = (
        "Ethics matter (Juerissen, 2000).\n\n## References\n\n"
        "Jeurissen, R. (2000) Cannibals with forks. Journal of Business Ethics."
    )
    fixed = cites.repair_citation_names(content)
    assert "(Jeurissen, 2000)" in fixed
    assert "Juerissen" not in fixed


def test_web_verification_drops_shop_links_and_cites_book_by_publisher(monkeypatch):
    payload = {
        "items": [
            {"cite": "Johnson 2017", "found": True, "title": "Exploring Strategy", "authors": ["Johnson, G."],
             "year": "2017", "publisher": "Pearson", "edition": "11th", "place": "Harlow",
             "url": "https://www.amazon.co.uk/Exploring-Strategy/dp/1292145129"},
            {"cite": "DESNZ 2024", "found": True, "title": None, "authors": ["Department for Energy Security"],
             "organisation": True, "year": "2024", "url": "https://www.gov.uk/x"},
        ]
    }
    seen: dict = {}

    def fake_search(**kw):
        seen.update(kw)
        return json.dumps(payload), [{"url": "https://www.gov.uk/x", "title": "Energy"}], {}

    monkeypatch.setattr("services.claude_client.generate_with_web_search", fake_search)
    confirmed = cites._web_verify_citations(
        ["Johnson 2017", "DESNZ 2024"],
        style="Harvard",
        contexts={"Johnson 2017": "Strategy is a long-term direction (Johnson et al., 2017)."},
    )
    assert [q for q, _ in confirmed] == ["Johnson 2017"]
    reference = confirmed[0][1]["reference"]
    assert "amazon" not in reference.lower() and "Accessed" not in reference
    assert "11th edn. Harlow: Pearson" in reference
    assert "long-term direction" in seen["user_prompt"]


def test_web_verification_accepts_url_whose_title_was_searched(monkeypatch):
    payload = {"items": [{"cite": "Ansoff 1957", "found": True, "title": "Strategies for Diversification",
                          "authors": ["Ansoff, H.I."], "year": "1957", "container": "Harvard Business Review",
                          "url": "https://hbr.org/1957/09/strategies-for-diversification"}]}
    sources = [{"url": "https://scholar.example.org/ansoff", "title": "Strategies for Diversification - HBR"}]
    monkeypatch.setattr(
        "services.claude_client.generate_with_web_search", lambda **_kw: (json.dumps(payload), sources, {})
    )
    confirmed = cites._web_verify_citations(["Ansoff 1957"], style="Harvard")
    assert [q for q, _ in confirmed] == ["Ansoff 1957"]


def test_engine_skips_topic_searches_when_the_writer_cited_enough(monkeypatch):
    engine = cites.AssignmentCitationEngine.__new__(cites.AssignmentCitationEngine)
    looked_up: list[str] = []

    def fake_lookup(query, **_kw):
        looked_up.append(query)
        return None

    engine._lookup = fake_lookup
    engine._verify_on_web = lambda *_a, **_kw: []
    engine._expand_used_citations = lambda body, **kw: (
        body, kw["works"], kw["references"], kw["in_text"], kw["queries"]
    )
    monkeypatch.setattr(cites, "_extract_queries", lambda *_a, **_kw: ["triple bottom line"])
    monkeypatch.setattr(cites, "_reconcile_body_with_works", lambda body, **_kw: body)
    draft = {"content": "## Intro\n\nA (Barney, 1991). B (Porter, 1980). C (Ansoff, 1957).", "version": 1}
    req = {"citation_style": "Harvard", "minimum_sources": 3}
    engine.generate(draft=draft, requirement_json=req, project_id="p")
    assert "triple bottom line" not in looked_up
    assert set(looked_up) == {"Barney 1991", "Porter 1980", "Ansoff 1957"}


# --- verified facts pack ------------------------------------------------------


def test_facts_pack_keeps_only_searched_official_figures():
    from services.assignment_facts import facts_as_works, facts_material, gather_verified_facts

    payload = {"facts": [
        {"fact": "CPI rose by 2.5% in the 12 months to December 2024.", "organisation": "Office for National Statistics",
         "year": "2024", "title": "Consumer price inflation, UK: December 2024",
         "url": "https://www.ons.gov.uk/cpi", "factor": "Economic"},
        {"fact": "Solar made 5.2% of generation.", "organisation": "Energy Saving Trust / UK Government",
         "year": "2025", "title": "Solar", "url": "https://www.ons.gov.uk/solar", "factor": "Environmental"},
        {"fact": "Invented 9% figure.", "organisation": "Somebody", "year": "2024", "title": "Made up",
         "url": "https://nowhere.example.com/x", "factor": "Social"},
    ]}
    sources = [{"url": "https://www.ons.gov.uk/cpi", "title": "CPI"}]
    facts = gather_verified_facts("PESTLE analysis of a UK firm", search=lambda **_kw: (json.dumps(payload), sources, {}))
    assert [f["organisation"] for f in facts] == ["Office for National Statistics"]
    works = facts_as_works(facts, style="Harvard")
    assert cites._work_matches_query("ONS 2024", works[0])
    assert "Office for National Statistics, 2024" in facts_material(facts)


def test_facts_pack_rejects_blogs_foreign_governments_and_misattribution():
    from services.assignment_facts import _published_by, gather_verified_facts

    assert _published_by("Office for National Statistics", "https://www.ons.gov.uk/x")
    assert _published_by("Department for Energy Security and Net Zero", "https://assets.publishing.service.gov.uk/x")
    assert not _published_by("Microgeneration Certification Scheme", "https://solarweekly.co.uk/uk-solar")
    assert _published_by("Microgeneration Certification Scheme", "https://mcscertified.com/data")
    payload = {"case_country": "United Kingdom", "facts": [
        {"fact": "EV charging market worth $2.1 billion in 2024.", "organisation": "International Trade Administration",
         "publisher_country": "United States", "year": "2024", "title": "UK EV charging",
         "url": "https://www.trade.gov/uk-ev", "factor": "Economic"},
    ]}
    sources = [{"url": "https://www.trade.gov/uk-ev", "title": "UK EV charging"}]
    assert gather_verified_facts("PESTLE", search=lambda **_kw: (json.dumps(payload), sources, {})) == []


def test_citations_prefer_pre_verified_facts():
    from services.assignment_facts import facts_as_works

    facts = [{"fact": "CPI rose 2.5%.", "organisation": "Office for National Statistics", "year": "2024",
              "title": "Consumer price inflation, UK: December 2024", "url": "https://www.ons.gov.uk/cpi"}]
    work = cites._best_verified_work("ONS 2024", facts_as_works(facts, style="Harvard"), context="Inflation fell.")
    assert work and work["title"].startswith("Consumer price inflation")


# --- writer coherence and trimming ------------------------------------------------


def test_summary_is_written_last_and_sees_the_report():
    from services.writer_engine.llm_writer import _build_section_prompt
    from services.writer_engine.models import WriterEngineInput, WriterSection, WriterSectionStatus
    from services.writer_engine.service import _active_section

    sections = [
        WriterSection(id="s1", title="Section 1: Executive Summary", objective="", estimated_words=200),
        WriterSection(id="s2", title="Section 2: Introduction", objective="", estimated_words=200),
        WriterSection(id="s3", title="Section 7: Conclusion", objective="", estimated_words=150),
    ]
    session = SimpleNamespace(sections=sections)
    assert _active_section(session).id == "s2"
    sections[1].status = WriterSectionStatus.COMPLETED
    sections[1].generated_text = "EcoNova faces a skills shortage."
    assert _active_section(session).id == "s3"
    sections[2].status = WriterSectionStatus.COMPLETED
    assert _active_section(session).id == "s1"

    payload = WriterEngineInput(
        requirement_json=_kaplan_requirement(),
        research_plan={"x": 1},
        blueprint={"sections": []},
        written_sections=[("Section 2: Introduction", "EcoNova faces a skills shortage.")],
    )
    prompt = _build_section_prompt(section=sections[0], payload=payload, revision=False)
    assert "REPORT_SO_FAR" in prompt and "skills shortage" in prompt
    assert "ONE STORY" in prompt and "inimitable" in prompt


def test_post_humanize_trim_removes_only_the_excess_and_keeps_section_edges():
    from services.writer_engine.gemini_trim import _just_enough, _Sentence

    sentences = [
        _Sentence(id=f"a-s{i}", section_title="A", text="one two three four five six seven eight nine ten",
                  index_in_section=i - 1)
        for i in range(1, 7)
    ]
    chosen = _just_enough(["a-s1", "a-s6", "a-s2", "a-s3", "a-s4"], sentences, need_words=15)
    assert chosen == ["a-s2", "a-s3"]


# --- formatting -----------------------------------------------------------------


def test_alignment_is_only_checked_when_the_brief_sets_it():
    spec = build_assignment_spec(_kaplan_requirement(), project_id="p")
    assert not spec.formatting.alignment_required
    req = _kaplan_requirement()
    req["formatting"] = {**req["formatting"], "alignment": "justified"}
    spec = build_assignment_spec(req, project_id="p")
    assert spec.formatting.alignment_required and spec.formatting.alignment == "justify"


def test_section_labels_restored_from_brief():
    from services.assignment_project.requirement_analyzer import _restore_section_labels

    brief = "Section 1: Executive Summary (Approx. 200 words)\nSection 8: References (not included)"
    out = _restore_section_labels(["Executive Summary", "References", "Appendix"], brief)
    assert out == ["Section 1: Executive Summary", "Section 8: References", "Appendix"]


def test_page_number_field_parts_are_separate_runs():
    from formatter_v2.render.styles import _add_field

    doc = docx.Document()
    paragraph = doc.add_paragraph()
    _add_field(paragraph, "PAGE")
    runs = paragraph._p.findall(docx.oxml.ns.qn("w:r"))
    assert len(runs) == 5
